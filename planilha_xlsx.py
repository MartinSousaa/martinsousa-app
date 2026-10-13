"""planilha_xlsx.py — abrir .xlsx lendo ATÉ A ÚLTIMA LINHA DE VERDADE.

13/10: o extrato do Itaú tinha 60 lançamentos e o Studio lia 2. O arquivo
declara no próprio cabeçalho o tamanho da tabela (`<dimension ref="A1:F13">`),
e o Itaú escreve esse tamanho errado. O `openpyxl` em modo somente leitura
(`read_only=True`, o rápido) CONFIA nessa declaração e para na linha 13 — sem
erro, sem aviso. Durante uma semana a culpa foi posta no banco ("exportou
cortado") e o dono exportou de novo à toa. Era o leitor.

`ws.reset_dimensions()` manda o leitor ignorar a declaração e ler até o fim.
Não custa nada quando a declaração está certa: ela só deixa de ser usada.

TODO .xlsx DO STUDIO ABRE POR AQUI. A conferência abaixo varre o repositório
e reprova `load_workbook(..., read_only=True)` fora deste arquivo: o defeito
existia em quatro leitores (extrato do Itaú, Controle MS, fatura do ML e o
gerador do histórico), e o quinto não pode nascer igual.
"""
import io


def abrir(origem, data_only=True):
    """O workbook em modo somente leitura, lendo todas as linhas.

    `origem` são os bytes do arquivo ou o caminho dele. Quem chama fecha.
    """
    import openpyxl
    fonte = io.BytesIO(origem) if isinstance(origem, (bytes, bytearray)) else origem
    wb = openpyxl.load_workbook(fonte, read_only=True, data_only=data_only)
    for ws in wb.worksheets:
        # Só a planilha somente-leitura tem a declaração para esquecer.
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
    return wb


def linhas_no_arquivo(dados, indice=0):
    """Quantas linhas a planilha `indice` TEM, contadas no XML do arquivo.

    Independe do leitor: é a medida de "o que a fonte tem" contra a qual se
    confere "o que o Studio leu" (Forma 10). None quando não dá para contar.
    """
    import re
    import zipfile
    try:
        z = zipfile.ZipFile(io.BytesIO(dados))
        nomes = sorted(n for n in z.namelist()
                       if re.match(r"xl/worksheets/sheet\d+\.xml$", n))
        if indice >= len(nomes):
            return None
        alvo = (f"xl/worksheets/sheet{indice + 1}.xml"
                if f"xl/worksheets/sheet{indice + 1}.xml" in nomes
                else nomes[indice])
        return len(re.findall(rb"<row[\s>]", z.read(alvo)))
    except Exception:
        return None


def _leitores_por_fora(raiz):
    """["arquivo:linha"] de `load_workbook(..., read_only=True)` fora daqui."""
    import ast
    import os
    fora = []
    for nome in sorted(os.listdir(raiz)):
        if (not nome.endswith(".py") or nome.startswith("checar_")
                or nome == "planilha_xlsx.py"):
            continue
        try:
            arv = ast.parse(open(os.path.join(raiz, nome),
                                 encoding="utf-8").read())
        except SyntaxError:
            continue
        for no in ast.walk(arv):
            if not isinstance(no, ast.Call):
                continue
            f = no.func
            nome_f = f.attr if isinstance(f, ast.Attribute) else getattr(f, "id", "")
            if nome_f != "load_workbook":
                continue
            if any(k.arg == "read_only" and isinstance(k.value, ast.Constant)
                   and k.value.value is True for k in no.keywords):
                fora.append(f"{nome}:{no.lineno}")
    return fora


# ── Conferência ──────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import os
    import re
    import sys
    import zipfile
    import openpyxl
    falhas = 0

    def ok(nome, cond):
        global falhas
        falhas += not cond
        print(("ok    " if cond else "FALHA ") + nome)

    # O ARQUIVO COM O DEFEITO DO ITAÚ: 95 linhas de verdade e o cabeçalho
    # declarando A1:F13 — o tamanho que o extrato de 13/10 declarava.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data", "Lançamento", "Razão Social", "CPF/CNPJ", "Valor (R$)",
               "Saldo (R$)"])
    for i in range(94):
        ws.append([f"{(i % 28) + 1:02d}/08/2026", f"PIX {i}", "X", "", -1.0, ""])
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                dados, _trocas = re.subn(rb'<dimension ref="[^"]+"\s*/>',
                                         b'<dimension ref="A1:F13"/>', dados)
                assert _trocas == 1, "o teste não conseguiu montar o defeito"
            zout.writestr(item, dados)
    defeituoso = out.getvalue()

    _cru = openpyxl.load_workbook(io.BytesIO(defeituoso), read_only=True,
                                  data_only=True)
    _n_cru = sum(1 for _ in _cru.worksheets[0].iter_rows(values_only=True))
    _cru.close()
    ok("o arquivo de teste tem o defeito do Itaú (o modo rápido lê só 13)",
       _n_cru == 13)
    _wb = abrir(defeituoso)
    _n = sum(1 for _ in _wb.worksheets[0].iter_rows(values_only=True))
    _wb.close()
    ok("abrir() lê as 95 linhas de verdade", _n == 95)

    # PELA CADEIA: o leitor do extrato do Itaú, sobre o mesmo arquivo.
    import extrato_itau as _ei
    _l, _cab, _err = _ei.ler(defeituoso)
    ok("o extrato do Itaú lê os 94 lançamentos, e não 12",
       _err == "" and len(_l) == 94)

    ok("conta as linhas no próprio arquivo, sem confiar na declaração",
       linhas_no_arquivo(defeituoso) == 95
       and linhas_no_arquivo(b"nao e xlsx") is None)
    ok("e o leitor do Itaú diz quantas não leu: nenhuma",
       _cab.get("_linhas_nao_lidas") == 0)

    # SE O DEFEITO VOLTAR, A TELA CULPA O STUDIO: o leitor parando na linha
    # 13 de novo (o modo rápido sem o reset) tem de contar o que não leu, e o
    # aviso da tela tem de dizer que a falta é do Studio, e não do banco.
    # Rodando como script, este arquivo é `__main__`; o leitor importa o
    # MÓDULO `planilha_xlsx` — é nele que se troca.
    import planilha_xlsx as _modulo
    _abrir_bom = _modulo.abrir
    try:
        _modulo.abrir = lambda origem, data_only=True: openpyxl.load_workbook(
            io.BytesIO(origem), read_only=True, data_only=data_only)
        _l2, _cab2, _ = _ei.ler(defeituoso)
    finally:
        _modulo.abrir = _abrir_bom
    import extratos_tela as _et
    _aviso = _et.extrato_incompleto("01/08/2026 até 13/10/2026", len(_l2),
                                    len(_l2) + _cab2["_linhas_nao_lidas"])
    ok("leitor parando no meio: conta as linhas que não leu",
       len(_l2) == 12 and _cab2["_linhas_nao_lidas"] == 82)
    ok("e a tela diz que a falta é do Studio", "falta é do Studio" in _aviso)

    _raiz = os.path.dirname(os.path.abspath(__file__))
    _fora = _leitores_por_fora(_raiz)
    ok("nenhum leitor de .xlsx abre no modo rápido por fora daqui",
       not _fora)
    if _fora:
        print("      " + ", ".join(_fora))
    print(f"\nfalhas: {falhas}")
    sys.exit(1 if falhas else 0)
