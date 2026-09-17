"""controle_ms.py — o Controle MS do Drive, lido sem ser convertido.

POR QUE ELE NÃO É UMA PLANILHA GOOGLE
-------------------------------------
O Studio inteiro fala com o Google Sheets (`sheets.py`). O Controle MS não é
uma planilha do Google: é um `.xlsx` de ~12 MB com tabela dinâmica, aberto no
modo Office do Sheets. `gspread` não abre isso — e a mensagem que ele devolve
faz parecer falta de permissão, o que manda consertar o que não está quebrado.

Converter resolveria a leitura e quebraria o trabalho: o dono alimenta essa
planilha todo dia, e no formato Google ele perderia a tabela dinâmica e o
Excel. Entre mudar o sistema e mudar o processo dele, o barato é o sistema —
o arquivo continua Excel, e quem se adapta é o Studio.

COMO
----
Baixa o arquivo pelo Drive (a conta de serviço precisa estar compartilhada
nele) e lê com openpyxl em modo `read_only`, que percorre as linhas sem montar
a planilha inteira na memória. Um container de 512 MB não aguenta um workbook
de 12 MB carregado por inteiro a cada clique.

O DOWNLOAD É CARO — E É POR ISSO QUE ELE É UM SÓ
------------------------------------------------
São 12 MB por leitura fria. Sem cache, cada rerun do Streamlit baixaria tudo
de novo: é o mesmo erro que já custou caro nos abonos e nas batidas, e a saída
é a mesma — ler uma vez e servir de memória, por TTL.

O QUE ELE NÃO FAZ
-----------------
Não escreve. O compartilhamento veio como Editor, mas escrever de volta num
arquivo que o dono edita à mão, ao mesmo tempo, é perder o que ele digitou —
o Drive não faz merge de `.xlsx`. Enquanto isso não for decidido, esta é uma
leitura e nada mais.
"""

import base64
import io
import re

import streamlit as st

# O arquivo mora no Drive novo da empresa. O ID não é segredo — é o endereço
# do arquivo, e ele aparece na URL para quem já tem acesso. O que é segredo é
# a chave da conta de serviço, e essa continua só nas Secrets.
ARQUIVO_ID_PADRAO = "1TX3kgzb815EekA3v5mg-9MRvSc-5Ggc3"

# Dez minutos: o dono edita a planilha em blocos, não célula a célula. Menos
# que isso paga 12 MB por um dado que não mudou.
TTL_SEG = 600

_MIME_XLSX = ("application/vnd.openxmlformats-officedocument"
              ".spreadsheetml.sheet")


def id_do_link(texto):
    """O ID do arquivo a partir da URL do Drive. "" quando não dá para achar.

    Existe porque o que chega pelo chat é o link inteiro, e pedir "só o ID"
    é transferir para o gestor um trabalho que uma expressão regular faz.
    """
    t = str(texto or "").strip()
    if not t:
        return ""
    m = re.search(r"/d/([a-zA-Z0-9_-]{15,})", t)
    if m:
        return m.group(1)
    m = re.search(r"[?&]id=([a-zA-Z0-9_-]{15,})", t)
    if m:
        return m.group(1)
    # Já veio só o ID.
    return t if re.fullmatch(r"[a-zA-Z0-9_-]{15,}", t) else ""


def arquivo_id():
    """O ID configurado. A secret manda; sem ela, vale o do código."""
    try:
        import gdrive as _gd
        da_secret = id_do_link(_gd._secret("CONTROLE_MS_ID", ""))
    except Exception:
        da_secret = ""
    return da_secret or ARQUIVO_ID_PADRAO


@st.cache_data(ttl=TTL_SEG, show_spinner=False)
def _baixar(file_id):
    """Os bytes do arquivo. (bytes, erro) — nunca levanta.

    `supportsAllDrives` vai SEMPRE: sem ele, um arquivo que mora numa Unidade
    Compartilhada responde "não encontrado" mesmo com o compartilhamento
    correto, e o diagnóstico manda conferir a permissão que já está certa.
    """
    try:
        import gdrive as _gd
        from googleapiclient.http import MediaIoBaseDownload
        svc = _gd.service()
        pedido = svc.files().get_media(fileId=file_id, supportsAllDrives=True)
        buf = io.BytesIO()
        baixador = MediaIoBaseDownload(buf, pedido, chunksize=4 * 1024 * 1024)
        fim = False
        while not fim:
            _status, fim = baixador.next_chunk()
        return buf.getvalue(), ""
    except Exception as e:
        try:
            import gdrive as _gd2
            return None, _gd2.erro_amigavel(e)
        except Exception:
            return None, str(e)[:200]



# ── De onde o arquivo vem ────────────────────────────────────────────────────
#
# Duas fontes possíveis, e a escolha não é técnica: é de onde o dono consegue
# trabalhar. Excel com salvamento automático e duas pessoas na mesma planilha
# só existe no OneDrive — é regra da Microsoft, não configuração. O Drive não
# tem isso e nunca vai ter.
#
# Então o arquivo fica onde o trabalho acontece, e o Studio vai buscá-lo lá:
# com um link de compartilhamento, baixando por HTTP puro, sem app da
# Microsoft e sem credencial nova.
#
# O link é SEGREDO, e por isso mora na secret CONTROLE_MS_URL e nunca no
# código: quem o tiver baixa a planilha financeira inteira, sem senha.
#
# Sem a secret, vale o arquivo do Google Drive — que continua funcionando e é
# o caminho sem link público.
SECRET_URL = "CONTROLE_MS_URL"

# Tempo de espera do download. Generoso: são 12 MB atravessando a internet, e
# conexão que não estabelece em 15s não vai estabelecer.
TIMEOUT = (15, 120)


def url_de_download(link):
    """O endereço que devolve os BYTES do arquivo. "" quando não dá.

    Link de compartilhamento abre uma PÁGINA, não o arquivo: baixar direto dele
    traz o HTML do visualizador, e o openpyxl recebe isso como se fosse a
    planilha — o erro que sai é "arquivo corrompido", que manda procurar defeito
    onde não há.

    As duas famílias de link precisam de tratamentos diferentes:

    - OneDrive pessoal (`1drv.ms`, `onedrive.live.com`): a própria Microsoft
      publica um atalho — o link em base64, com prefixo `u!`, vira um endereço
      de conteúdo na API de compartilhamentos. Não exige autenticação nenhuma,
      porque o link já é a autorização.
    - OneDrive corporativo / SharePoint: aceita `download=1` na própria URL.
    """
    t = str(link or "").strip()
    if not t.lower().startswith("http"):
        return ""
    baixo = t.lower()
    # Já é um endereço de conteúdo: não mexer.
    if "/root/content" in baixo or "download=1" in baixo:
        return t
    if "sharepoint.com" in baixo:
        return t + ("&" if "?" in t else "?") + "download=1"
    if "1drv.ms" in baixo or "onedrive.live.com" in baixo:
        b64 = base64.urlsafe_b64encode(t.encode("utf-8")).decode("ascii")
        return f"https://api.onedrive.com/v1.0/shares/u!{b64.rstrip('=')}/root/content"
    # Desconhecido: devolve como veio. Se for um link direto, funciona; se não
    # for, o erro aparece na tela de conferência com o endereço na mão.
    return t


def link_configurado():
    """O link do OneDrive guardado na secret. "" quando não há."""
    try:
        import gdrive as _gd
        return str(_gd._secret(SECRET_URL, "") or "").strip()
    except Exception:
        return ""


def fonte():
    """('onedrive'|'drive', descrição curta) — de onde o arquivo está vindo.

    A tela precisa dizer isto em voz alta: com duas fontes possíveis, "o número
    está velho" tem duas causas diferentes, e saber qual está ativa é a
    primeira pergunta.
    """
    if link_configurado():
        return "onedrive", "Link do OneDrive (secret CONTROLE_MS_URL)"
    return "drive", f"Google Drive · {arquivo_id()}"


@st.cache_data(ttl=TTL_SEG, show_spinner=False)
def _baixar_do_link(url):
    """Os bytes do arquivo pelo link público. (bytes, erro)."""
    import requests
    alvo = url_de_download(url)
    if not alvo:
        return None, "O link configurado em CONTROLE_MS_URL não é um endereço."
    try:
        r = requests.get(alvo, timeout=TIMEOUT, allow_redirects=True)
    except Exception as e:
        return None, f"O OneDrive não respondeu: {type(e).__name__}: {str(e)[:120]}"
    if not r.ok:
        return None, (f"O OneDrive respondeu HTTP {r.status_code}. O link pode "
                      f"ter sido revogado ou trocado — gere outro e atualize a "
                      f"secret {SECRET_URL}.")
    dados = r.content
    # Link de PÁGINA devolve HTML com status 200. Sem esta checagem, o erro só
    # aparece lá adiante, como "arquivo corrompido".
    if dados[:2] != b"PK":
        return None, ("O link devolveu uma página, e não o arquivo. Use a opção "
                      "de compartilhamento que gera link de download, ou "
                      "confira se o acesso é «qualquer pessoa com o link».")
    return dados, ""


def _bytes_do_arquivo():
    """(bytes, erro) da fonte ativa."""
    link = link_configurado()
    if link:
        return _baixar_do_link(link)
    return _baixar(arquivo_id())


def _workbook(dados):
    """O workbook em modo somente leitura. Quem chama fecha."""
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(dados), read_only=True,
                                  data_only=True)


def abas():
    """(nomes_das_abas, erro). Lista vazia quando não deu para ler."""
    dados, erro = _bytes_do_arquivo()
    if erro or not dados:
        return [], erro or "A fonte não devolveu o arquivo."
    try:
        wb = _workbook(dados)
        try:
            return list(wb.sheetnames), ""
        finally:
            wb.close()
    except Exception as e:
        return [], f"Não consegui abrir o arquivo: {type(e).__name__}: {str(e)[:140]}"


def ler(aba, limite=None):
    """(DataFrame, erro) de uma aba. A primeira linha é o cabeçalho.

    `limite` corta o número de linhas de dados — é o que a tela de conferência
    usa para mostrar as primeiras sem pagar a planilha inteira.
    """
    import pandas as pd
    dados, erro = _bytes_do_arquivo()
    if erro or not dados:
        return pd.DataFrame(), erro or "A fonte não devolveu o arquivo."
    try:
        wb = _workbook(dados)
    except Exception as e:
        return pd.DataFrame(), f"Não consegui abrir o arquivo: {str(e)[:140]}"
    try:
        if aba not in wb.sheetnames:
            return pd.DataFrame(), (f"A aba «{aba}» não existe. Existem: "
                                    + ", ".join(wb.sheetnames))
        ws = wb[aba]
        linhas = []
        cabecalho = None
        for linha in ws.iter_rows(values_only=True):
            if cabecalho is None:
                cabecalho = _cabecalho(linha)
                continue
            if all(c is None or str(c).strip() == "" for c in linha):
                continue
            linhas.append(linha[:len(cabecalho)])
            if limite and len(linhas) >= limite:
                break
        return pd.DataFrame(linhas, columns=cabecalho or []), ""
    except Exception as e:
        return pd.DataFrame(), f"Falhou lendo «{aba}»: {str(e)[:140]}"
    finally:
        wb.close()


def _cabecalho(linha):
    """Nomes de coluna utilizáveis a partir da primeira linha.

    Coluna sem nome vira `coluna_3`, e nome repetido ganha sufixo: o pandas
    aceita nomes duplicados e depois devolve duas colunas no mesmo `df[nome]`,
    o que transforma uma soma em erro silencioso.
    """
    fora, vistos = [], {}
    for i, c in enumerate(linha or [], start=1):
        nome = str(c).strip() if c is not None and str(c).strip() else f"coluna_{i}"
        if nome in vistos:
            vistos[nome] += 1
            nome = f"{nome}_{vistos[nome]}"
        else:
            vistos[nome] = 1
        fora.append(nome)
    return fora


def limpar():
    """Esquece o arquivo baixado. Para a tela que acabou de pedir releitura."""
    for f in (_baixar, _baixar_do_link):
        try:
            f.clear()
        except Exception:
            pass


# ── Conferência ──────────────────────────────────────────────────────────────
# Só a parte que decide. O download depende do Drive e de credencial; o que se
# confere aqui é o que erra em silêncio: ID lido do link e cabeçalho repetido.
if __name__ == "__main__":
    falhas = 0

    def ok(nome, cond):
        global falhas
        falhas += not cond
        print(("ok    " if cond else "FALHA ") + nome)

    _ID = "1TX3kgzb815EekA3v5mg-9MRvSc-5Ggc3"
    ok("o ID sai do link inteiro, com gid e tudo",
       id_do_link(f"https://docs.google.com/spreadsheets/d/{_ID}/edit?gid=49#gid=49")
       == _ID)
    ok("link sem /edit também", id_do_link(f"https://drive.google.com/file/d/{_ID}/view") == _ID)
    ok("formato antigo, com ?id=", id_do_link(f"https://drive.google.com/open?id={_ID}") == _ID)
    ok("o próprio ID passa direto", id_do_link(_ID) == _ID)
    ok("texto que não é link devolve vazio",
       id_do_link("manda o link depois") == "" and id_do_link("") == ""
       and id_do_link(None) == "")

    ok("cabeçalho normal fica como está",
       _cabecalho(("Data", "Valor")) == ["Data", "Valor"])
    # Coluna repetida e coluna sem nome: as duas viram erro silencioso no
    # pandas, uma somando duas colunas e a outra sumindo da tela.
    ok("nome repetido ganha sufixo, e não some",
       _cabecalho(("Valor", "Valor", "Valor")) == ["Valor", "Valor_2", "Valor_3"])
    ok("coluna sem nome vira posição",
       _cabecalho(("Data", None, "  ")) == ["Data", "coluna_2", "coluna_3"])
    ok("linha vazia não derruba", _cabecalho(()) == [] and _cabecalho(None) == [])

    # ── O link de compartilhamento vira endereço de ARQUIVO ───────────────
    # Sem isto, o download traz a pagina do visualizador e o openpyxl acusa
    # "arquivo corrompido" — mandando procurar defeito onde nao ha.
    _pess = "https://1drv.ms/x/s!AbCdEf123456"
    _u = url_de_download(_pess)
    ok("link do OneDrive pessoal vira endereco de conteudo",
       _u.startswith("https://api.onedrive.com/v1.0/shares/u!")
       and _u.endswith("/root/content"))
    ok("o base64 do link nao leva '=' no fim (a API recusa)",
       "=" not in _u.split("u!")[1].split("/")[0])
    ok("o link original esta dentro, e da para voltar dele",
       base64.urlsafe_b64decode(
           _u.split("u!")[1].split("/")[0] + "==").decode() == _pess)

    _corp = "https://empresa-my.sharepoint.com/:x:/g/personal/leo/Documento.xlsx"
    ok("SharePoint ganha download=1",
       url_de_download(_corp) == _corp + "?download=1")
    ok("SharePoint com query usa &",
       url_de_download(_corp + "?e=abc") == _corp + "?e=abc&download=1")
    ok("endereco que ja baixa nao e mexido",
       url_de_download(_corp + "?download=1") == _corp + "?download=1")
    ok("texto que nao e endereco devolve vazio",
       url_de_download("manda depois") == "" and url_de_download("") == ""
       and url_de_download(None) == "")

    print("\nfalhas:", falhas)
