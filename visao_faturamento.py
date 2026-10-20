"""visao_faturamento.py — a aba Faturamento da Visão Geral.

O PEDIDO, NAS PALAVRAS DO DONO (15/10)
--------------------------------------
A aba DINAMICA do Controle MS e o resumo que ele passa ao time todo mês, dentro
do Studio: faturado, vendas, UC e saldo UC; a meta do ano; os trimestres contra
a meta; o UC mês a mês; o mês a mês contra 2025; os marketplaces; e um resumo
pronto para recortar. Mais duas coisas que a planilha não faz:

  - o faturamento do mês corrente "vai atualizando com base no Bling";
  - a CONFERÊNCIA pedido a pedido: "quais vendas estão na minha planilha e não
    estão com ele, ou quais estão com ele computadas do Bling e não estão na
    minha planilha. Ele tem que me passar a informação completa."

DE ONDE VEM CADA NÚMERO
-----------------------
    meses fechados        BASE DE VENDAS (FAT - DEV, VENDAS, L.B, DEVOLUÇÃO)
    mês corrente          Bling, `home_gestao._faturamento_bling` — o mesmo da Home
    2025 (comparação)     FAT. CONSOLIDADO
    metas                 DINAMICA (meta do ano e a de cada trimestre); sem a
                          DINAMICA do ano, os trimestres do MELHOR ano anterior
                          (dono: "a meta trimestral precisa ser com base no
                          nosso melhor ano")
    UC                    a coluna UNI. CONT. da planilha: MARGEM C. ÷ LPV, por
                          venda. NÃO é peça por venda. No mês sem LPV ainda, o
                          último LPV fechado.

A CONFERÊNCIA SÓ RODA EM MÊS PRONTO
-----------------------------------
Mês pronto = toda venda com SKU e preço de custo. Antes disso a planilha está
pela metade, e comparar acusaria divergência que não existe.
"""
import io
import math
import re
from datetime import date, datetime, timedelta

import streamlit as st

# Formato de dinheiro e nome do mês: os mesmos da aba Financeira, e não cópia
# (Forma 5 — duas respostas para a mesma pergunta passam a discordar).
from visao_geral import _MES_TXT, _brl  # noqa: E402

MESES = ("JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT",
         "NOV", "DEZ")
ABA_BASE = "BASE DE VENDAS 2026"
ABA_CONSOLIDADO = "FAT. CONSOLIDADO"
ABA_DINAMICA = "DINAMICA"
DIAS_UC = 30
TOLERANCIA_PEDIDO = 0.50


def _num(v):
    try:
        f = float(v)
        return f if f == f and not math.isinf(f) else 0.0
    except (TypeError, ValueError):
        try:
            return float(str(v).replace(".", "").replace(",", "."))
        except (TypeError, ValueError):
            return 0.0


def _vazio(v):
    return str(v if v is not None else "").strip() in ("", "nan", "None", "0",
                                                       "0.0", "NaT")


def chave_pedido(v):
    """O número do pedido para comparar planilha com Bling. Função pura.

    A planilha guarda o pedido do ML como texto ('2000015134692443) e às vezes
    como número (2.000015134692443e15 no pandas); o Bling manda texto. Só
    letras e dígitos, em maiúscula.
    """
    if isinstance(v, float) and v == v and v.is_integer():
        v = str(int(v))
    return re.sub(r"[^0-9A-Z]", "", str(v or "").upper())


def _data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()[:10]
    try:
        return date.fromisoformat(t)
    except ValueError:
        return None


def mes_da_linha(l):
    """(ano, mes) da venda: pela Data; sem ela, pelo MÊS/ANO ('jan2026')."""
    d = _data(l.get("Data"))
    if d:
        return d.year, d.month
    m = re.match(r"([a-zç]{3})\D*(\d{4})", str(l.get("MÊS/ANO") or "").lower())
    if m and m.group(1) in _MES_TXT:
        return int(m.group(2)), _MES_TXT.index(m.group(1)) + 1
    return None, None


def lpv_atual(linhas):
    """O LPV do último mês que já tem LPV na planilha (coluna LPV)."""
    melhor = None
    for l in linhas or []:
        lpv = _num(l.get("LPV"))
        if lpv <= 0:
            continue
        a, m = mes_da_linha(l)
        if a and (melhor is None or (a, m) > melhor[0]):
            melhor = ((a, m), lpv)
    return melhor[1] if melhor else None


def uc_da_linha(l, lpv_reserva):
    """A UC da venda: a coluna UNI. CONT.; vazia, a mesma fórmula da planilha
    (SE(DEVOLUÇÃO>0;"";SE(FAT TOTAL=0;"";MARGEM C. ÷ LPV))) com o LPV da linha
    ou, sem ele, o último LPV fechado."""
    u = _num(l.get("UNI. CONT."))
    if u:
        return u
    if _num(l.get("DEVOLUÇÃO")) > 0 or _num(l.get("FAT TOTAL")) == 0:
        return 0.0
    lpv = _num(l.get("LPV")) or float(lpv_reserva or 0)
    return _num(l.get("MARGEM C.")) / lpv if lpv else 0.0


def venda_da_linha(l):
    """1 quando a linha é venda: a coluna VENDAS; vazia, a mesma fórmula da
    planilha (SE(FAT TOTAL=0;"";SE(DEVOLUÇÃO>0;"";1))). Sem isto, a linha nova
    de outubro — ainda sem a fórmula recalculada — entrava na UC e não nas
    vendas, e a UC dos 30 dias saía 1,52 em vez de 1,42."""
    v = _num(l.get("VENDAS"))
    if v:
        return v
    return 1.0 if _num(l.get("FAT TOTAL")) and not _num(l.get("DEVOLUÇÃO")) > 0 else 0.0


def resumo_mensal(linhas, ano, lpv_reserva):
    """{mes: {...}} da BASE DE VENDAS no ano. Função pura."""
    fora = {}
    for l in linhas or []:
        a, m = mes_da_linha(l)
        if a != int(ano):
            continue
        r = fora.setdefault(m, {"fat_total": 0.0, "fat_liq": 0.0, "dev": 0.0,
                                "lb": 0.0, "vendas": 0.0, "uc_total": 0.0,
                                "linhas": 0, "sem_sku": 0, "sem_custo": 0})
        r["linhas"] += 1
        r["fat_total"] += _num(l.get("FAT TOTAL"))
        r["fat_liq"] += _num(l.get("FAT - DEV"))
        r["dev"] += _num(l.get("DEVOLUÇÃO"))
        r["lb"] += _num(l.get("L.B"))
        r["vendas"] += venda_da_linha(l)
        r["uc_total"] += uc_da_linha(l, lpv_reserva)
        r["sem_sku"] += _vazio(l.get("SKU"))
        r["sem_custo"] += _vazio(l.get("Custo Unitário"))
    return fora


def pronto(r):
    """Mês pronto para conferir: toda venda com SKU e preço de custo."""
    return bool(r) and r["linhas"] > 0 and not r["sem_sku"] and not r["sem_custo"]


def uc_ultimos_dias(linhas, lpv_reserva, dias=DIAS_UC):
    """(uc_medio, saldo, inicio, fim) das vendas dos últimos `dias` da planilha."""
    datas = [d for d in (_data(l.get("Data")) for l in linhas or []) if d]
    if not datas:
        return None, None, None, None
    fim = max(datas)
    ini = fim - timedelta(days=dias - 1)
    uc = vd = 0.0
    for l in linhas:
        d = _data(l.get("Data"))
        if d and ini <= d <= fim:
            uc += uc_da_linha(l, lpv_reserva)
            vd += venda_da_linha(l)
    return (uc / vd if vd else None), round(uc - vd), ini, fim


def por_plataforma(linhas, ano):
    """{plataforma: FAT - DEV} no ano, do maior para o menor."""
    fora = {}
    for l in linhas or []:
        if mes_da_linha(l)[0] == int(ano):
            p = str(l.get("Plataforma") or "sem plataforma").strip()
            fora[p] = fora.get(p, 0.0) + _num(l.get("FAT - DEV"))
    return dict(sorted(fora.items(), key=lambda kv: -kv[1]))


def consolidado(linhas):
    """{ano: {mes: (faturamento, vendas)}} da aba FAT. CONSOLIDADO.

    A aba tem dois blocos lado a lado (FAT e VENDAS), cada um com uma coluna
    por ano; `controle_ms.ler` dá sufixo `_2` ao ano repetido do segundo.
    """
    fora, vistos = {}, set()
    for l in linhas or []:
        rot = str(l.get("FAT") or "").strip().upper()[:3]
        if rot not in MESES:
            continue
        m = MESES.index(rot) + 1
        # SÓ O PRIMEIRO BLOCO: mais abaixo a aba repete JAN…DEZ com os % de
        # crescimento, e a segunda passada trocava R$ 260.952 por 1,63.
        if m in vistos:
            continue
        vistos.add(m)
        for col, v in l.items():
            c = str(col)
            if re.fullmatch(r"20\d\d", c):
                fora.setdefault(int(c), {}).setdefault(m, [0.0, 0.0])[0] = _num(v)
            elif re.fullmatch(r"20\d\d_2", c):
                fora.setdefault(int(c[:4]), {}).setdefault(m, [0.0, 0.0])[1] = _num(v)
    return {a: {m: tuple(v) for m, v in ms.items()} for a, ms in fora.items()}


def metas_da_dinamica(celulas, ano):
    """{anual, fat: [4], vendas: [4]} lidos da DINAMICA, ou None. Função pura.

    `celulas` = [(linha, coluna, valor)]. A meta do ano fica à direita do rótulo
    "META <ano>"; cada trimestre aparece duas vezes ("1º Trimestre" … ): o
    primeiro bloco é de faturamento, o segundo de vendas, e a meta do período
    fica duas colunas à direita do rótulo (P → R).
    """
    pos = {(r, c): v for r, c, v in celulas or []}
    ano_aba = next((v for (r, c), v in sorted(pos.items())
                    if isinstance(v, (int, float)) and 2020 <= v <= 2100), None)
    if ano_aba is not None and int(ano_aba) != int(ano):
        return None
    anual = None
    for (r, c), v in sorted(pos.items()):
        if str(v).strip().upper() == f"META {int(ano)}":
            x = pos.get((r, c + 1))
            if isinstance(x, (int, float)) and x > 1000:
                anual = float(x)
                break
    blocos = {}
    for (r, c), v in sorted(pos.items()):
        mt = re.fullmatch(r"([1-4])º\s*TRIMESTRE", str(v).strip().upper())
        if mt:
            x = pos.get((r, c + 2))
            if isinstance(x, (int, float)) and x > 0:
                blocos.setdefault(int(mt.group(1)), []).append(float(x))
    if not all(len(blocos.get(q, [])) >= 2 for q in (1, 2, 3, 4)):
        return None
    return {"anual": anual, "fat": [blocos[q][0] for q in (1, 2, 3, 4)],
            "vendas": [blocos[q][1] for q in (1, 2, 3, 4)], "fonte": "DINAMICA"}


def metas_do_melhor_ano(cons, ano):
    """Os trimestres do melhor ano anterior, como meta. Função pura."""
    anos = {a: ms for a, ms in (cons or {}).items() if a < int(ano)
            and len(ms) == 12 and all(v[0] for v in ms.values())}
    if not anos:
        return None
    melhor = max(anos, key=lambda a: sum(v[0] for v in anos[a].values()))
    ms = anos[melhor]
    fat = [sum(ms[m][0] for m in range(q * 3 + 1, q * 3 + 4)) for q in range(4)]
    ven = [sum(ms[m][1] for m in range(q * 3 + 1, q * 3 + 4)) for q in range(4)]
    return {"anual": sum(fat), "fat": fat, "vendas": ven,
            "fonte": f"trimestres de {melhor}, o melhor ano"}


# ── Conferência Bling × planilha ─────────────────────────────────────────────

def pedidos_da_planilha(linhas, ano, mes):
    """{chave: {...}} — a venda de vários produtos (várias linhas) vira UMA."""
    fora = {}
    for l in linhas or []:
        if mes_da_linha(l) != (int(ano), int(mes)):
            continue
        k = chave_pedido(l.get("Pedido"))
        if not k:
            continue
        p = fora.setdefault(k, {"pedido": str(l.get("Pedido") or "").strip(),
                                "data": (_data(l.get("Data")) or "") and _data(l.get("Data")).isoformat(),
                                "plataforma": str(l.get("Plataforma") or "").strip(),
                                "produto": [], "sku": [], "valor": 0.0})
        p["valor"] = round(p["valor"] + _num(l.get("FAT TOTAL")), 2)
        if str(l.get("Produto") or "").strip():
            p["produto"].append(str(l.get("Produto")).strip()[:60])
        if str(l.get("SKU") or "").strip():
            p["sku"].append(str(l.get("SKU")).strip())
    for p in fora.values():
        p["produto"] = " + ".join(p["produto"])
        p["sku"] = ", ".join(p["sku"])
    return fora


def conferir(planilha, bling, ano, mes):
    """As três listas e o resumo. Função pura.

    `planilha` = `pedidos_da_planilha` do mês; `bling` = pedidos do Bling de UM
    POUCO ANTES a UM POUCO DEPOIS do mês (uma venda do dia 31 pode cair no dia
    1º do outro lado). "Só no Bling" conta só o pedido com data dentro do mês.
    """
    alvo = f"{int(ano):04d}-{int(mes):02d}"
    bl = {}
    for b in bling or []:
        k = chave_pedido(b.get("pedido")) or chave_pedido(b.get("numero_bling"))
        if k:
            x = bl.setdefault(k, dict(b, total=0.0))
            x["total"] = round(x["total"] + float(b.get("total") or 0), 2)
            x["cancelado"] = x.get("cancelado") and b.get("cancelado")
    so_planilha = [dict(p, onde="só na planilha") for k, p in planilha.items()
                   if k not in bl]
    so_bling = [dict(b, onde="só no Bling") for k, b in bl.items()
                if k not in planilha and str(b.get("data", "")).startswith(alvo)]
    diferentes = []
    for k, p in planilha.items():
        b = bl.get(k)
        if b and abs(p["valor"] - b["total"]) > TOLERANCIA_PEDIDO:
            diferentes.append({"pedido": p["pedido"], "data": p["data"],
                               "plataforma": p["plataforma"],
                               "planilha": p["valor"], "bling": b["total"],
                               "diferenca": round(p["valor"] - b["total"], 2),
                               "cancelado_no_bling": bool(b.get("cancelado"))})
    no_mes = {k: b for k, b in bl.items() if str(b.get("data", "")).startswith(alvo)
              or k in planilha}
    resumo = {"pedidos_planilha": len(planilha), "pedidos_bling": len(no_mes),
              "valor_planilha": round(sum(p["valor"] for p in planilha.values()), 2),
              "valor_bling": round(sum(b["total"] for b in no_mes.values()
                                       if not b.get("cancelado")), 2),
              "cancelado_bling": round(sum(b["total"] for b in no_mes.values()
                                           if b.get("cancelado")), 2)}
    return {"so_planilha": sorted(so_planilha, key=lambda x: x["data"]),
            "so_bling": sorted(so_bling, key=lambda x: x["data"]),
            "diferentes": sorted(diferentes, key=lambda x: -abs(x["diferenca"])),
            "resumo": resumo}


def planilha_da_conferencia(conf):
    """Os bytes do .xlsx com as três listas, para baixar."""
    import openpyxl
    wb = openpyxl.Workbook()
    abas = (("Só na planilha", conf["so_planilha"],
             ("pedido", "data", "plataforma", "produto", "sku", "valor")),
            ("Só no Bling", conf["so_bling"],
             ("pedido", "numero_bling", "data", "canal", "total", "cancelado")),
            ("Valor diferente", conf["diferentes"],
             ("pedido", "data", "plataforma", "planilha", "bling", "diferenca",
              "cancelado_no_bling")))
    for i, (nome, linhas, cols) in enumerate(abas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = nome
        ws.append(list(cols))
        for l in linhas:
            ws.append([l.get(c) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Leitura ──────────────────────────────────────────────────────────────────

@st.cache_data(ttl=600, show_spinner=False)
def ler_planilha():
    """({base, consolidado, dinamica}, avisos). Uma leitura a cada 10 min."""
    import controle_ms as _cms
    import planilha_xlsx as _px
    avisos, fora = [], {}
    for nome, aba in (("base", ABA_BASE), ("consolidado", ABA_CONSOLIDADO)):
        df, erro = _cms.ler(aba)
        if erro:
            avisos.append(f"{aba}: {erro}")
        fora[nome] = df.to_dict("records") if not erro and df is not None else []
    celulas = []
    try:
        dados, erro = _cms._bytes_do_arquivo()
        if dados and not erro:
            wb = _px.abrir(dados)
            try:
                if ABA_DINAMICA in wb.sheetnames:
                    for r, linha in enumerate(wb[ABA_DINAMICA].iter_rows(values_only=True), 1):
                        if r > 40:
                            break
                        for c, v in enumerate(linha, 1):
                            if v is not None and str(v).strip():
                                celulas.append((r, c, v))
            finally:
                wb.close()
    except Exception as e:
        avisos.append(f"{ABA_DINAMICA}: {type(e).__name__}")
    fora["dinamica"] = celulas
    return fora, avisos


def bling_do_mes_corrente(hoje):
    """(faturamento líquido do mês no Bling, avisos) — a mesma conta da Home."""
    import home_gestao as _hg
    return _hg._faturamento_bling(hoje.year, hoje.month)


class _FalhaBling(Exception):
    def __init__(self, lista, erro):
        super().__init__(erro)
        self.lista, self.erro = lista, erro


@st.cache_data(ttl=1800, show_spinner=False)
def _pedidos_da_conta(conta, ini, fim):
    """Só o que veio SEM erro fica guardado: o `st.cache_data` não guarda
    exceção, então uma falha do Bling é tentada de novo no próximo clique,
    em vez de ficar meia hora na tela."""
    import bling_api as _bl
    lista, erro = _bl.pedidos_resumidos(conta, ini, fim)
    if erro:
        raise _FalhaBling(lista, erro)
    return lista


def pedidos_do_bling(ano, mes):
    """(pedidos das duas contas de 3 dias antes a 3 dias depois do mês, erros)."""
    import calendar
    import bling_api as _bl
    ini = date(int(ano), int(mes), 1) - timedelta(days=3)
    fim = date(int(ano), int(mes), calendar.monthrange(int(ano), int(mes))[1]) + timedelta(days=3)
    todos, erros = [], []
    for conta in _bl.CONTAS:
        try:
            todos += _pedidos_da_conta(conta, ini.isoformat(), fim.isoformat())
        except _FalhaBling as f:
            erros.append(f"{conta.upper()}: {f.erro}")
            todos += f.lista
    return todos, erros


def dados_do_ano(ano, hoje, plan):
    """Tudo o que a aba desenha, de `ler_planilha`. Pura no que importa."""
    base = plan.get("base") or []
    lpv_r = lpv_atual(base)
    meses = resumo_mensal(base, ano, lpv_r)
    cons = consolidado(plan.get("consolidado"))
    metas = metas_da_dinamica(plan.get("dinamica"), ano) or metas_do_melhor_ano(cons, ano)
    uc30 = uc_ultimos_dias(base, lpv_r)
    return {"meses": meses, "anterior": cons.get(int(ano) - 1, {}), "metas": metas,
            "uc30": uc30, "lpv_reserva": lpv_r,
            "plataformas": por_plataforma(base, ano)}


# ── Tela ─────────────────────────────────────────────────────────────────────

FUNDO, CARTAO, BORDA, TILE = "#3c3c3c", "#333333", "#4a4a4a", "#2c2c2c"
TEXTO, SEC, MUTED = "#E8EAED", "#9AA3AE", "#6b737c"
REAL, ML, SHP, SHE = "#2E9E6A", "#4F93D6", "#C2801C", "#9579E0"
BOM, RUIM = "#4CC38A", "#E5534B"
VERDE_A, AZUL_A = "#5be0a0", "#7db4ec"
VIDRO = (f"background:linear-gradient(160deg,#3a3a3a 0%,#2e2e2e 100%);"
         f"border:1px solid {BORDA};border-radius:16px;"
         f"box-shadow:0 8px 24px rgba(0,0,0,.25)")


def _n(v):
    return "—" if v is None else f"{float(v):,.0f}".replace(",", ".")


def _p(v, sinal=True):
    if v is None:
        return "—"
    return (("+" if v > 0 and sinal else "") + f"{v:.1f}%").replace(".", ",")


def _k(v):
    return "—" if v is None else f"R$ {v / 1000:.0f}k"


def _uc(v):
    return "—" if v is None else f"{v:.2f}".replace(".", ",")


def _vidro(t, miolo):
    return (f'<div style="{VIDRO};padding:18px;margin-top:14px">'
            f'<div style="text-align:center;font-size:12px;font-weight:800;'
            f'letter-spacing:.14em;color:{SEC};margin-bottom:14px">{t}</div>{miolo}</div>')


def _chip(pc):
    if pc is None:
        return ""
    bom = pc >= 0
    return (f'<span style="display:inline-block;padding:2px 8px;border-radius:999px;'
            f'font-size:11px;font-weight:700;background:{"#1f4a36" if bom else "#4a2321"};'
            f'color:{BOM if bom else RUIM}">{"▲" if bom else "▼"} {_p(pc)}</span>')


def realizado_por_mes(d, hoje, ano, bling_mes):
    """{mes: faturamento} — fechado pela planilha, o corrente pelo Bling."""
    fora = {}
    for m, r in d["meses"].items():
        fora[m] = r["fat_liq"]
    if int(ano) == hoje.year and bling_mes is not None:
        fora[hoje.month] = bling_mes
    return fora


def trimestres(d, fat_mes):
    """[(q, meta_fat, real_fat, meta_vendas, real_vendas, uc, saldo)]."""
    mt = d["metas"] or {"fat": [None] * 4, "vendas": [None] * 4}
    fora = []
    for q in range(4):
        ms = range(q * 3 + 1, q * 3 + 4)
        com = [m for m in ms if m in d["meses"]]
        fechado = len(com) == 3
        rf = sum(fat_mes.get(m, 0.0) for m in ms) if fechado else None
        rv = sum(d["meses"][m]["vendas"] for m in com) if fechado else None
        ut = sum(d["meses"][m]["uc_total"] for m in com)
        vt = sum(d["meses"][m]["vendas"] for m in com)
        fora.append((q + 1, mt["fat"][q], rf, mt["vendas"][q], rv,
                     (ut / vt) if fechado and vt else None,
                     round(ut - vt) if fechado else None))
    return fora


def _svg_donut(pct):
    raio, w = 78, 14
    cf = 2 * math.pi * raio
    return (f'<svg viewBox="0 0 200 200" width="210" height="210"><defs>'
            f'<linearGradient id="gD" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="{VERDE_A}"/>'
            f'<stop offset="1" stop-color="#1f8f5c"/></linearGradient></defs>'
            f'<circle cx="100" cy="100" r="{raio}" fill="none" stroke="#2a2a2a" stroke-width="{w}"/>'
            f'<circle cx="100" cy="100" r="{raio}" fill="none" stroke="url(#gD)" stroke-width="{w}" '
            f'stroke-linecap="round" stroke-dasharray="{min(pct, 100) / 100 * cf:.1f} {cf:.1f}" '
            f'transform="rotate(-90 100 100)"/>'
            f'<text x="100" y="98" text-anchor="middle" fill="{TEXTO}" font-size="36" font-weight="800">{_p(pct, False)}</text>'
            f'<text x="100" y="120" text-anchor="middle" fill="{SEC}" font-size="11" letter-spacing="2">DA META</text></svg>')


def _svg_acumulado(fat_mes, metas_tri):
    W, H, L, B_ = 560, 200, 52, 24
    meta_mes = [q / 3 for q in metas_tri for _ in range(3)] if metas_tri else []
    acm, s = [], 0.0
    for m in range(1, 13):
        if m not in fat_mes:
            break
        s += fat_mes[m]
        acm.append(s)
    mac, s = [], 0.0
    for v in meta_mes:
        s += v
        mac.append(s)
    mx = max(mac + acm + [1.0]) * 1.05
    X = lambda i: L + (i + 0.5) * (W - L) / 12
    Y = lambda v: H - B_ - v / mx * (H - B_ - 10)
    fm = lambda v: (f"R$ {v / 1e6:.2f} mi".replace(".", ",") if v >= 1e6 else _k(v))
    s_ = "".join(f'<rect x="{X(q * 3) - (W - L) / 24:.1f}" y="10" width="{(W - L) / 4:.1f}" height="{H - B_ - 10}" fill="#fff" fill-opacity="{0.035 if q % 2 else 0}"/>'
                 f'<text x="{X(q * 3 + 1):.1f}" y="22" text-anchor="middle" fill="{MUTED}" font-size="10">{q + 1}º TRI</text>' for q in range(4))
    s_ += "".join(f'<line x1="{L}" x2="{W}" y1="{Y(mx * f):.1f}" y2="{Y(mx * f):.1f}" stroke="{BORDA}"/>'
                  f'<text x="{L - 6}" y="{Y(mx * f) + 3:.1f}" text-anchor="end" fill="{MUTED}" font-size="10">{fm(mx * f)}</text>'
                  for f in (0.25, 0.5, 0.75, 1.0))
    if mac:
        s_ += (f'<polyline points="{" ".join(f"{X(i):.1f},{Y(v):.1f}" for i, v in enumerate(mac))}" '
               f'fill="none" stroke="{TEXTO}" stroke-width="2" stroke-dasharray="5 5" stroke-opacity=".7"/>')
    if acm:
        pts = " ".join(f"{X(i):.1f},{Y(v):.1f}" for i, v in enumerate(acm))
        u = len(acm) - 1
        s_ += (f'<polygon points="{X(0):.1f},{Y(0):.1f} {pts} {X(u):.1f},{Y(0):.1f}" fill="{REAL}" fill-opacity=".15"/>'
               f'<polyline points="{pts}" fill="none" stroke="{REAL}" stroke-width="2.5"/>'
               f'<circle cx="{X(u):.1f}" cy="{Y(acm[u]):.1f}" r="5" fill="{REAL}"/>'
               f'<text x="{X(u) + 8:.1f}" y="{Y(acm[u]) + 16:.1f}" fill="{TEXTO}" font-size="11" font-weight="700">{fm(acm[u])}</text>')
    s_ += "".join(f'<text x="{X(i):.1f}" y="{H - 6}" text-anchor="middle" fill="{MUTED}" font-size="10">{m}</text>' for i, m in enumerate(MESES))
    leg = (f'<div style="font-size:11px;color:{SEC}"><span style="color:{REAL}">━</span> realizado acumulado'
           f' &nbsp; <span style="color:{TEXTO}">╌</span> meta acumulada</div>')
    return f'<b>Faturamento acumulado</b><svg viewBox="0 0 {W} {H}" width="100%">{s_}</svg>{leg}'


def _barra_meta(rot, meta, real, fmt, unid=""):
    if meta is None:
        return ""
    if real is None:
        return (f'<div style="margin-top:10px;font-size:12px;display:flex;justify-content:space-between">'
                f'<span style="color:{SEC};font-weight:700">{rot}</span><span style="color:{MUTED}">meta {fmt(meta)} · em andamento</span></div>')
    mx = max(meta, real) * 1.1
    return (f'<div style="margin-top:10px"><div style="display:flex;justify-content:space-between;font-size:12px">'
            f'<span style="color:{SEC};font-weight:700">{rot}</span><span><b>{fmt(real)}</b> <span style="color:{MUTED}">/ meta {fmt(meta)}</span></span></div>'
            f'<div style="height:9px;background:{TILE};border-radius:5px;margin-top:5px;position:relative">'
            f'<div style="width:{real / mx * 100:.1f}%;height:9px;background:{REAL};border-radius:5px"></div>'
            f'<div style="position:absolute;left:{meta / mx * 100:.1f}%;top:-3px;width:2px;height:15px;background:{TEXTO}"></div></div>'
            f'<div style="margin-top:5px">{_chip((real / meta - 1) * 100)}</div></div>')


def _svg_uc(d, hoje, ano):
    meses = d["meses"]
    ucm = {m: (r["uc_total"] / r["vendas"]) for m, r in meses.items() if r["vendas"]}
    q4 = any(m in ucm for m in (10, 11, 12)) and len([m for m in (10, 11, 12) if m in ucm]) == 3
    n_mes = 12 if q4 else 9
    larg30 = 0 if q4 else 1.6
    W, H, L, B_ = 1120, 250, 46, 26
    un = (W - L) / (n_mes + larg30)
    lo, hi = 0.80, 1.70
    X = lambda i: L + (i + 0.5) * un
    Y = lambda v: H - B_ - (min(max(v, lo), hi) - lo) / (hi - lo) * (H - B_ - 44)
    s = ""
    for q in range(n_mes // 3):
        x0 = L + q * 3 * un
        ms = [m for m in range(q * 3 + 1, q * 3 + 4) if m in meses]
        ut, vt = sum(meses[m]["uc_total"] for m in ms), sum(meses[m]["vendas"] for m in ms)
        uq = ut / vt if vt and len(ms) == 3 else None
        s += f'<rect x="{x0:.1f}" y="10" width="{3 * un:.1f}" height="{H - B_ - 10}" fill="#fff" fill-opacity="{0.035 if q % 2 else 0}"/>'
        s += (f'<text x="{X(q * 3 + 1):.1f}" y="27" text-anchor="middle" fill="{TEXTO if uq else MUTED}" font-size="12" '
              f'font-weight="700">{q + 1}º TRI · {("UC " + _uc(uq)) if uq else "em andamento"}</text>')
        if uq:
            s += (f'<line x1="{x0 + 10:.1f}" x2="{x0 + 3 * un - 10:.1f}" y1="{Y(uq):.1f}" y2="{Y(uq):.1f}" '
                  f'stroke="{REAL}" stroke-dasharray="4 6" stroke-opacity=".55"/>')
    for f in (1.0, 1.2, 1.4, 1.6):
        s += (f'<line x1="{L}" x2="{W}" y1="{Y(f):.1f}" y2="{Y(f):.1f}" stroke="{BORDA}" stroke-opacity=".5"/>'
              f'<text x="{L - 8}" y="{Y(f) + 4:.1f}" text-anchor="end" fill="{MUTED}" font-size="11">{_uc(f)}</text>')
    P = [(X(m - 1), Y(ucm[m]), ucm[m]) for m in range(1, n_mes + 1) if m in ucm]
    if len(P) >= 2:
        curva = f"M{P[0][0]:.1f},{P[0][1]:.1f}"
        for i in range(len(P) - 1):
            p0 = P[i - 1] if i else P[i]
            p1, p2 = P[i], P[i + 1]
            p3 = P[i + 2] if i + 2 < len(P) else p2
            curva += (f" C{p1[0] + (p2[0] - p0[0]) / 6:.1f},{p1[1] + (p2[1] - p0[1]) / 6:.1f} "
                      f"{p2[0] - (p3[0] - p1[0]) / 6:.1f},{p2[1] - (p3[1] - p1[1]) / 6:.1f} {p2[0]:.1f},{p2[1]:.1f}")
        base_y = Y(lo)
        s += f'<path d="{curva} L{P[-1][0]:.1f},{base_y:.1f} L{P[0][0]:.1f},{base_y:.1f} Z" fill="{REAL}" fill-opacity=".10"/>'
        s += f'<path d="{curva}" fill="none" stroke="{REAL}" stroke-width="2" stroke-linecap="round"/>'
    for x, y, v in P:
        s += (f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.5" fill="{REAL}" stroke="#333" stroke-width="2"/>'
              f'<text x="{x:.1f}" y="{y - 10:.1f}" text-anchor="middle" fill="{TEXTO}" font-size="11" font-weight="700" '
              f'stroke="#323232" stroke-width="4" paint-order="stroke">{_uc(v)}</text>')
    s += "".join(f'<text x="{X(i):.1f}" y="{H - 7}" text-anchor="middle" fill="{MUTED}" font-size="11">{MESES[i]}</text>' for i in range(n_mes))
    uc30, sal30, i30, f30 = d["uc30"]
    if not q4 and uc30:
        x0 = L + n_mes * un + 8
        w = larg30 * un - 8
        cx = x0 + w / 2
        s += (f'<rect x="{x0:.1f}" y="10" width="{w:.1f}" height="{H - B_ - 10}" rx="10" fill="#fff" fill-opacity=".05" stroke="{BORDA}"/>'
              f'<text x="{cx:.1f}" y="27" text-anchor="middle" fill="{SEC}" font-size="11" font-weight="700">ÚLTIMOS {DIAS_UC} DIAS</text>'
              f'<text x="{cx:.1f}" y="{Y(uc30) + 6:.1f}" text-anchor="middle" fill="{TEXTO}" font-size="30" font-weight="800">{_uc(uc30)}</text>'
              f'<text x="{cx:.1f}" y="{Y(uc30) + 26:.1f}" text-anchor="middle" fill="{SEC}" font-size="11">saldo {_n(sal30)}</text>'
              f'<text x="{cx:.1f}" y="{H - 7}" text-anchor="middle" fill="{MUTED}" font-size="10">{i30:%d/%m} a {f30:%d/%m}</text>')
    tot_uc = sum(r["uc_total"] for r in meses.values())
    tot_v = sum(r["vendas"] for r in meses.values())
    cab = (f'<div style="display:flex;justify-content:space-between;align-items:baseline"><b>UC mês a mês</b>'
           f'<span style="font-size:12px;color:{SEC}">UC no ano <b style="color:{TEXTO}">{_uc(tot_uc / tot_v if tot_v else None)}</b>'
           f' · saldo <b style="color:{TEXTO}">{_n(tot_uc - tot_v)}</b></span></div>')
    leg = (f'<div style="font-size:11px;color:{MUTED}"><span style="color:{REAL}">━</span> UC do mês &nbsp; '
           f'<span style="color:{REAL}">╌</span> fechamento do trimestre &nbsp;·&nbsp; UC = margem de contribuição ÷ LPV</div>')
    return cab + f'<svg viewBox="0 0 {W} {H}" width="100%">{s}</svg>' + leg


def _tabela(d, fat_mes, hoje, ano, bling=False):
    # CENTRALIZADO (dono, 20/10: "esses números devem estar centralizados").
    th = lambda c: (f'<th style="text-align:center;padding:7px 8px;color:{SEC};font-size:11px;font-weight:700;'
                    f'border-bottom:1px solid {BORDA};white-space:nowrap">{c}</th>')
    td = lambda v, c=TEXTO, b=False: (f'<td style="text-align:center;padding:6px 8px;font-size:12px;white-space:nowrap;'
                                      f'color:{c};{"font-weight:800;" if b else ""}">{v}</td>')
    ant = d["anterior"]
    cols = (f"FAT {int(ano) - 1}", f"FAT {ano}", "CRESC.", f"VENDAS {int(ano) - 1}", f"VENDAS {ano}", "CRESC.",
            "TICKET MÉDIO", "DEVOLUÇÃO", "DEV. %", "FAT. LÍQUIDO", "MARGEM L.B", "UC TOTAL", "UC", "SALDO UC")
    corpo = ""
    for m in range(1, 13):
        r = d["meses"].get(m)
        f0, v0 = ant.get(m, (None, None))
        zebra = "background:rgba(255,255,255,.025);" if m % 2 == 0 else ""
        corrente = int(ano) == hoje.year and m == hoje.month
        if r:
            f1 = fat_mes.get(m)
            # mês corrente é parcial: contra o mês cheio do ano anterior o
            # crescimento sairia sempre negativo — fica "—" até fechar.
            cf = ((f1 / f0 - 1) * 100) if f0 and f1 and not corrente else None
            cv = ((r["vendas"] / v0 - 1) * 100) if v0 and r["vendas"] and not corrente else None
            seta = lambda x: "" if x is None else ("▲ " if x >= 0 else "▼ ")
            cor = lambda x: SEC if x is None else (BOM if x >= 0 else RUIM)
            corpo += (f'<tr style="{zebra}"><td style="padding:6px 8px;font-weight:700">{MESES[m - 1]}'
                      f'{" <span style=font-size:9px;color:#7db4ec>BLING</span>" if corrente and bling else ""}</td>'
                      + td(_brl(f0), SEC) + td(_brl(f1)) + td(seta(cf) + _p(cf), cor(cf), True)
                      + td(_n(v0), SEC) + td(_n(r["vendas"])) + td(seta(cv) + _p(cv), cor(cv), True)
                      + td(_brl(r["fat_total"] / r["vendas"], 2) if r["vendas"] else "—")
                      + td(_brl(r["dev"])) + td(_p(r["dev"] / r["fat_total"] * 100, False) if r["fat_total"] else "—")
                      + td(_brl(r["fat_liq"])) + td(_p(r["lb"] / r["fat_liq"] * 100, False) if r["fat_liq"] else "—")
                      + td(_n(r["uc_total"])) + td(_uc(r["uc_total"] / r["vendas"]) if r["vendas"] else "—", TEXTO, True)
                      + td(_n(r["uc_total"] - r["vendas"])) + "</tr>")
        else:
            corpo += (f'<tr style="{zebra}"><td style="padding:6px 8px;font-weight:700">{MESES[m - 1]}</td>'
                      + td(_brl(f0), SEC) + td("—", MUTED) * 2 + td(_n(v0), SEC) + td("—", MUTED) * 10 + "</tr>")
    return (f'<div style="overflow-x:auto"><table style="border-collapse:collapse;width:100%">'
            f'<tr><th></th>{"".join(th(c) for c in cols)}</tr>{corpo}</table></div>')


_COR_PLAT = {"mercadolivre": ML, "shopee": SHP, "shein": SHE}


def _cor_plat(nome, idx):
    base = next((c for k, c in _COR_PLAT.items() if k in nome.lower().replace(" ", "")), MUTED)
    return base, max(1 - 0.18 * idx, 0.35)


def _plataformas(plats):
    T = sum(plats.values()) or 1.0
    vistos, itens = {}, []
    for nome, v in plats.items():
        base, _ = _cor_plat(nome, 0)
        i = vistos.get(base, 0)
        vistos[base] = i + 1
        itens.append((nome, v, base, max(1 - 0.18 * i, 0.35)))
    ang, arcs = -90.0, ""
    for nome, v, c, op in itens:
        a2 = ang + v / T * 360
        x1, y1 = 100 + 80 * math.cos(math.radians(ang)), 100 + 80 * math.sin(math.radians(ang))
        x2, y2 = 100 + 80 * math.cos(math.radians(a2 - 0.8)), 100 + 80 * math.sin(math.radians(a2 - 0.8))
        arcs += (f'<path d="M{x1:.1f},{y1:.1f} A80,80 0 {1 if a2 - ang > 180 else 0},1 {x2:.1f},{y2:.1f}" '
                 f'fill="none" stroke="{c}" stroke-opacity="{op}" stroke-width="26"/>')
        ang = a2
    rosca = (f'<svg viewBox="0 0 200 200" width="180" height="180">{arcs}'
             f'<text x="100" y="96" text-anchor="middle" fill="{TEXTO}" font-size="17" font-weight="800">{_brl(T)}</text>'
             f'<text x="100" y="116" text-anchor="middle" fill="{SEC}" font-size="10">{len(itens)} contas</text></svg>')
    lista = "".join(f'<div style="display:flex;justify-content:space-between;padding:6px 0;border-bottom:1px solid {BORDA};font-size:13px">'
                    f'<span><span style="display:inline-block;width:10px;height:10px;border-radius:3px;background:{c};opacity:{op};margin-right:8px"></span>{nome}</span>'
                    f'<span><b>{_brl(v)}</b> <span style="color:{MUTED}">· {_p(v / T * 100, False)}</span></span></div>'
                    for nome, v, c, op in itens)
    return (f'<div style="display:grid;grid-template-columns:190px 1fr;gap:16px;align-items:center">'
            f'<div style="text-align:center">{rosca}</div>'
            f'<div>{lista}</div></div>')


def _top_produtos(top, n_total=5):
    """Os 5 produtos que mais venderam no ano, com a UC e o peso no
    faturamento (dono, 20/10: "do lado direito, um top 5 produtos mais
    vendidos, a UC dele e o quanto representa do faturamento anual")."""
    if not top:
        return (f'<div style="color:{MUTED};font-size:12px;text-align:center;padding:30px 0">'
                f'Top 5 indisponível: nenhuma venda do ano foi lida agora.</div>')
    cor_uc = lambda u: BOM if u >= 1 else ("#E0A13A" if u >= 0.8 else RUIM)
    linhas = "".join(
        f'<div style="display:grid;grid-template-columns:22px 1fr auto auto;gap:10px;align-items:center;'
        f'padding:7px 0;border-bottom:1px solid {BORDA};font-size:13px">'
        f'<span style="color:{SEC};font-weight:800">{i}</span>'
        f'<span style="white-space:nowrap;overflow:hidden;text-overflow:ellipsis">{nome.title()}'
        f'<span style="color:{MUTED};font-size:11px"> · {_n(v)} vendas</span></span>'
        f'<span style="font-weight:800;color:{cor_uc(uc)}">UC {_uc(uc)}</span>'
        f'<span style="color:{SEC};min-width:56px;text-align:right">{_p(pct, False)}</span></div>'
        for i, (nome, v, uc, pct) in enumerate(top[:n_total], 1))
    return (f'<div style="font-size:11px;font-weight:800;letter-spacing:.1em;color:{SEC};margin-bottom:4px">'
            f'TOP 5 MAIS VENDIDOS · UC · % DO FATURAMENTO DO ANO</div>{linhas}')


def top_do_ano(ano, hoje):
    """[(produto, vendas, uc, % do faturamento)] — os títulos juntados por
    `produtos_grupos`, a conta de `produtos_analise`. [] se não der para ler."""
    try:
        import produtos_analise as _pa
        import produtos_grupos as _pg
        brutas, erro = _pa.dados(int(ano), hoje)
        if erro or not brutas:
            return []
        mapa, _ = _pg.carregar()
        linhas = _pg.aplicar(brutas, mapa)
        return _pa.top_vendidos(linhas)
    except Exception:
        return []


def _trilha(rot, real, meta, fmt, cor):
    if meta is None or real is None:
        return (f'<div style="margin:10px 0;font-size:13px;display:flex;justify-content:space-between">'
                f'<b>{rot}</b><span style="color:{MUTED}">{("meta " + fmt(meta)) if meta else ""} · em andamento</span></div>')
    pc = real / meta * 100
    return (f'<div style="margin:10px 0"><div style="display:flex;justify-content:space-between;align-items:baseline">'
            f'<span style="font-size:13px;font-weight:700">{rot}</span><span style="font-size:13px"><b>{fmt(real)}</b> '
            f'<span style="color:{MUTED}">de {fmt(meta)}</span> &nbsp;{_chip(pc - 100)}</span></div>'
            f'<div style="position:relative;height:12px;background:#2a2a2a;border-radius:6px;margin-top:6px">'
            f'<div style="height:12px;width:{min(pc, 115) / 1.15:.1f}%;background:linear-gradient(90deg,{cor}88,{cor});border-radius:6px"></div>'
            f'<div style="position:absolute;left:{100 / 1.15:.1f}%;top:-4px;width:2px;height:20px;background:{TEXTO}"></div></div></div>')


def html_resumo_time(d, tri, fat_ano, vendas_ano, uc_ano, saldo_ano, meta_anual, ate):
    """O resumo para o time (opção 2 aprovada): o círculo da meta e as trilhas."""
    pct = fat_ano / meta_anual * 100 if meta_anual else 0
    kp = lambda a, b, c: (f'<div style="{VIDRO};padding:14px"><div style="font-size:11px;color:{SEC};letter-spacing:.12em">{a}</div>'
                          f'<div style="font-size:26px;font-weight:800">{b}</div><div style="font-size:11px">{c}</div></div>')
    mf = sum(t[1] for t in tri if t[2] is not None and t[1])
    rf = sum(t[2] for t in tri if t[2] is not None)
    mv = sum(t[3] for t in tri if t[4] is not None and t[3])
    rv = sum(t[4] for t in tri if t[4] is not None)
    blocos = lambda tit, im, ir, fmt, cor: (
        f'<div style="{VIDRO};padding:16px"><div style="font-size:11px;font-weight:800;letter-spacing:.14em;color:{SEC}">{tit}</div>'
        + "".join(_trilha(f"{t[0]}º tri", t[ir], t[im], fmt, cor) for t in tri)
        + f'<div style="font-size:10px;color:{MUTED};margin-top:4px">traço branco = 100% da meta</div></div>')
    ucs = [(f"{t[0]}º tri", t[5]) for t in tri if t[5]] + ([(f"{DIAS_UC} dias", d["uc30"][0])] if d["uc30"][0] else [])
    uc_lins = "".join(
        f'<div style="display:grid;grid-template-columns:70px 1fr 60px;gap:10px;align-items:center;margin:10px 0">'
        f'<span style="font-size:13px;font-weight:700">{nome}</span><div style="height:12px;background:#2a2a2a;border-radius:6px">'
        f'<div style="height:12px;width:{min(uc / 1.7, 1) * 100:.0f}%;background:'
        f'{"repeating-linear-gradient(45deg," + VERDE_A + "99 0 4px,#2a3a33 4px 8px)" if "dias" in nome else "linear-gradient(90deg," + VERDE_A + "88," + VERDE_A + ")"};'
        f'border-radius:6px"></div></div><span style="text-align:right;font-weight:800">{_uc(uc)}</span></div>' for nome, uc in ucs)
    return (f'<div style="background:radial-gradient(1200px 400px at 20% 0%,#454545 0%,{FUNDO} 60%);border-radius:18px;padding:22px;border:1px solid {BORDA}">'
            f'<div style="display:flex;justify-content:space-between;align-items:flex-end;margin-bottom:16px">'
            f'<div><div style="font-size:22px;font-weight:300;letter-spacing:.42em">MARTINSOUSA</div>'
            f'<div style="font-size:12px;color:{SEC};letter-spacing:.14em;margin-top:2px">RESULTADO · ATÉ {ate}</div></div>'
            f'<div style="text-align:right;font-size:12px;color:{SEC}">Meta do ano<div style="font-size:20px;font-weight:800;color:{TEXTO}">{_brl(meta_anual)}</div></div></div>'
            f'<div style="display:grid;grid-template-columns:250px 1fr;gap:12px">'
            f'<div style="{VIDRO};padding:10px;display:flex;flex-direction:column;align-items:center;justify-content:center">{_svg_donut(pct)}'
            f'<div style="font-size:12px;color:{SEC}">faltam <b style="color:{TEXTO}">{_brl(max((meta_anual or 0) - fat_ano, 0))}</b></div></div>'
            f'<div><div style="display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:12px">'
            + kp("FATURADO", _brl(fat_ano), _chip((rf / mf - 1) * 100) if mf else "")
            + kp("VENDAS", _n(vendas_ano), _chip((rv / mv - 1) * 100) if mv else "")
            + kp("UC", _uc(uc_ano), f'<span style="color:{SEC}">margem ÷ LPV</span>')
            + kp("SALDO UC", _n(saldo_ano), f'<span style="color:{SEC}">UC total − vendas</span>')
            + '</div><div style="display:grid;grid-template-columns:repeat(3,1fr);gap:12px">'
            + blocos("FATURAMENTO · TRIMESTRES", 1, 2, _k, VERDE_A) + blocos("VENDAS · TRIMESTRES", 3, 4, _n, AZUL_A)
            + f'<div style="{VIDRO};padding:16px"><div style="font-size:11px;font-weight:800;letter-spacing:.14em;color:{SEC}">UC · TRIMESTRES</div>{uc_lins}</div>'
            + '</div></div></div></div>')


def _tabela_conf(linhas, cols, rotulos):
    import pandas as pd
    return pd.DataFrame([{rotulos[i]: l.get(c) for i, c in enumerate(cols)} for l in linhas])


def pagina(usuario_logado, ano, hoje):
    plan, avisos = ler_planilha()
    d = dados_do_ano(int(ano), hoje, plan)
    bling_mes, av_b = (None, [])
    if int(ano) == hoje.year:
        bling_mes, av_b = bling_do_mes_corrente(hoje)
    fat_mes = realizado_por_mes(d, hoje, ano, bling_mes)
    tri = trimestres(d, fat_mes)
    metas = d["metas"] or {}
    fat_ano = sum(fat_mes.values())
    vendas_ano = sum(r["vendas"] for r in d["meses"].values())
    uc_tot = sum(r["uc_total"] for r in d["meses"].values())
    uc_ano = uc_tot / vendas_ano if vendas_ano else None
    saldo_ano = uc_tot - vendas_ano
    mf = sum(t[1] for t in tri if t[2] is not None and t[1])
    rf = sum(t[2] for t in tri if t[2] is not None)
    mv = sum(t[3] for t in tri if t[4] is not None and t[3])
    rv = sum(t[4] for t in tri if t[4] is not None)
    card = lambda a, b, c: (f'<div style="{VIDRO};padding:14px;text-align:center"><div style="font-size:12px;color:{SEC}">{a}</div>'
                            f'<div style="font-size:26px;font-weight:800;margin:2px 0">{b}</div><div style="font-size:11px;color:{MUTED}">{c}</div></div>')
    acum = lambda tit, txt, bom: (f'<div style="{VIDRO};padding:12px;text-align:center"><span style="color:{SEC};font-size:12px">{tit}</span>'
                                  f'<div style="font-size:22px;font-weight:800;color:{BOM if bom else RUIM}">{txt}</div></div>')
    topo = (f'<div style="display:grid;grid-template-columns:repeat(4,1fr);gap:12px">'
            + card("Faturado no ano", _brl(fat_ano, 2), "faturamento − devolução"
                   + (" · mês corrente pelo Bling" if bling_mes is not None else ""))
            + card("Vendas no ano", _n(vendas_ano), "pedidos")
            + card("UC", _uc(uc_ano), "margem de contribuição ÷ LPV")
            + card("Saldo UC", _n(saldo_ano), "UC total − vendas") + '</div>')
    if mf or mv:
        topo += (f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:12px;margin-top:12px">'
                 + (acum("FATURAMENTO contra a meta (trimestres fechados)",
                         f'{"▲" if rf >= mf else "▼"} {_p((rf / mf - 1) * 100)} · {_brl(rf - mf, 2)}', rf >= mf) if mf else "<div></div>")
                 + (acum("VENDAS contra a meta (trimestres fechados)",
                         f'{"▲" if rv >= mv else "▼"} {_p((rv / mv - 1) * 100)} · {_n(rv - mv)}', rv >= mv) if mv else "<div></div>")
                 + '</div>')
    st.markdown(topo, unsafe_allow_html=True)
    meta_anual = metas.get("anual")
    if meta_anual:
        pct = fat_ano / meta_anual * 100
        meta_box = (f'<div style="display:flex;flex-direction:column;align-items:center">{_svg_donut(pct)}'
                    f'<div style="font-size:18px;font-weight:800">Meta {ano} · {_brl(meta_anual)}</div>'
                    f'<div style="font-size:13px;color:{SEC}">faltam <b style="color:{TEXTO}">{_brl(max(meta_anual - fat_ano, 0))}</b></div>'
                    f'<div style="font-size:11px;color:{MUTED};margin-top:4px">meta: {metas.get("fonte")}</div></div>')
    else:
        meta_box = f'<div style="color:{SEC};text-align:center">Sem meta do ano na DINAMICA nem ano anterior completo.</div>'
    st.markdown(_vidro(f"META {ano} E O ANO ACUMULADO",
                       f'<div style="display:grid;grid-template-columns:320px 1fr;gap:24px;align-items:center">'
                       f'{meta_box}<div>{_svg_acumulado(fat_mes, metas.get("fat"))}</div></div>'), unsafe_allow_html=True)
    cartoes = "".join(
        f'<div style="{VIDRO};padding:14px"><div style="text-align:center;font-weight:800;letter-spacing:.1em">{q}º TRIMESTRE</div>'
        + _barra_meta("Faturamento", mf_, rf_, _k) + _barra_meta("Vendas", mv_, rv_, _n)
        + (f'<div style="margin-top:12px;display:flex;justify-content:space-between;border-top:1px solid {BORDA};padding-top:10px">'
           f'<span style="color:{SEC};font-size:12px;font-weight:700">UC</span><span style="font-size:18px;font-weight:800">{_uc(uc_)}</span></div>'
           f'<div style="display:flex;justify-content:space-between;font-size:12px"><span style="color:{SEC}">saldo UC</span><b>{_n(sal_)}</b></div>'
           if uc_ else f'<div style="margin-top:12px;border-top:1px solid {BORDA};padding-top:10px;color:{MUTED};font-size:12px">UC: em andamento</div>')
        + '</div>' for q, mf_, rf_, mv_, rv_, uc_, sal_ in tri)
    st.markdown(_vidro("RESULTADOS TRIMESTRAIS", f'<div style="display:grid;grid-template-columns:repeat(4,1fr);gap:12px">{cartoes}</div>'),
                unsafe_allow_html=True)
    st.markdown(_vidro("UC · RESULTADOS MENSAIS E TRIMESTRAIS", _svg_uc(d, hoje, ano)), unsafe_allow_html=True)
    st.markdown(_vidro(f"MÊS A MÊS · {int(ano) - 1} × {ano}", _tabela(d, fat_mes, hoje, ano, bling_mes is not None)), unsafe_allow_html=True)
    if d["plataformas"]:
        # METADE E METADE (dono, 20/10): plataformas à esquerda, os 5 mais
        # vendidos à direita — o quadro tinha muito espaço para pouca coisa.
        st.markdown(_vidro(f"FATURAMENTO POR MARKETPLACE · {ano}",
                           f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:28px;align-items:start">'
                           f'<div>{_plataformas(d["plataformas"])}</div>'
                           f'<div>{_top_produtos(top_do_ano(ano, hoje))}</div></div>'), unsafe_allow_html=True)

    # ── CONFERÊNCIA BLING × PLANILHA ──
    st.markdown(_vidro("CONFERÊNCIA · BLING × PLANILHAS DAS PLATAFORMAS",
                       f'<div style="font-size:12px;color:{SEC};text-align:center">Só mês pronto: toda venda com SKU e preço de custo. '
                       f'Compara pedido a pedido pelo número da plataforma.</div>'), unsafe_allow_html=True)
    estado = []
    for m in sorted(d["meses"]):
        r = d["meses"][m]
        estado.append((m, pronto(r), r["sem_sku"], r["sem_custo"]))
    prontos = [m for m, ok, _, _ in estado if ok]
    pend = [f"{MESES[m - 1]}: {max(s1, s2)} venda(s) sem SKU/custo" for m, ok, s1, s2 in estado if not ok]
    if pend:
        st.caption("Aguardando a planilha · " + " · ".join(pend))
    if prontos:
        mes_c = st.radio("Mês para conferir", [MESES[m - 1] for m in prontos], index=len(prontos) - 1,
                         horizontal=True, key="vf_conf_mes", label_visibility="collapsed")
        mc = MESES.index(mes_c) + 1
        if st.button(f"🔎 Conferir {mes_c}/{ano} com o Bling", key="vf_conf_btn"):
            st.session_state["vf_conf_feito"] = (int(ano), mc)
        if st.session_state.get("vf_conf_feito") == (int(ano), mc):
            with st.spinner("Lendo os pedidos do Bling…"):
                bl, err = pedidos_do_bling(int(ano), mc)
            for e in err:
                st.warning(f"Bling {e}")
            pl = pedidos_da_planilha(plan.get("base"), int(ano), mc)
            conf = conferir(pl, bl, int(ano), mc)
            rs = conf["resumo"]
            if bl and not any(b.get("pedido") for b in bl):
                st.warning("O Bling respondeu, mas sem o número do pedido da plataforma (numeroLoja): "
                           "a comparação pedido a pedido não tem como casar. Me avise com um print.")
            dif = rs["valor_planilha"] - rs["valor_bling"]
            st.markdown(
                f'<div style="display:grid;grid-template-columns:repeat(4,1fr);gap:12px">'
                + card("Planilha", _brl(rs["valor_planilha"], 2), f'{_n(rs["pedidos_planilha"])} pedidos')
                + card("Bling (sem cancelados)", _brl(rs["valor_bling"], 2), f'{_n(rs["pedidos_bling"])} pedidos · cancelado {_brl(rs["cancelado_bling"], 2)}')
                + card("Diferença", _brl(dif, 2), "planilha − Bling")
                + card("Pedidos a olhar", _n(len(conf["so_planilha"]) + len(conf["so_bling"]) + len(conf["diferentes"])),
                       f'{len(conf["so_planilha"])} só na planilha · {len(conf["so_bling"])} só no Bling · {len(conf["diferentes"])} com valor diferente')
                + '</div>', unsafe_allow_html=True)
            if conf["so_planilha"]:
                st.markdown("**Na planilha e não no Bling**")
                st.dataframe(_tabela_conf(conf["so_planilha"], ("pedido", "data", "plataforma", "produto", "sku", "valor"),
                                          ("Pedido", "Data", "Plataforma", "Produto", "SKU", "Valor")),
                             hide_index=True, use_container_width=True)
            if conf["so_bling"]:
                st.markdown("**No Bling e não na planilha**")
                st.dataframe(_tabela_conf(conf["so_bling"], ("pedido", "numero_bling", "data", "canal", "total", "cancelado"),
                                          ("Pedido", "Nº no Bling", "Data", "Canal", "Valor", "Cancelado no Bling")),
                             hide_index=True, use_container_width=True)
            if conf["diferentes"]:
                st.markdown("**Nos dois, com valor diferente**")
                st.dataframe(_tabela_conf(conf["diferentes"], ("pedido", "data", "plataforma", "planilha", "bling", "diferenca", "cancelado_no_bling"),
                                          ("Pedido", "Data", "Plataforma", "Planilha", "Bling", "Diferença", "Cancelado no Bling")),
                             hide_index=True, use_container_width=True)
            if not (conf["so_planilha"] or conf["so_bling"] or conf["diferentes"]):
                st.success(f"✓ {mes_c}/{ano} bate pedido a pedido.")
            st.download_button("⬇ Baixar as listas (Excel)", planilha_da_conferencia(conf),
                               file_name=f"conferencia_bling_{ano}_{mc:02d}.xlsx", key="vf_conf_xlsx")

    # ── RESUMO PARA O TIME ──
    ult = max(d["meses"]) if d["meses"] else hoje.month
    resumo = html_resumo_time(d, tri, fat_ano, vendas_ano, uc_ano, saldo_ano, meta_anual, f"{MESES[ult - 1]}/{ano}")
    st.markdown(_vidro("RESUMO PARA O TIME · PRONTO PARA RECORTAR", resumo), unsafe_allow_html=True)
    pagina_html = (f'<!doctype html><html><head><meta charset="utf-8"><title>Resumo {ano}</title></head>'
                   f'<body style="background:{FUNDO};color:{TEXTO};font-family:-apple-system,Segoe UI,Roboto,sans-serif;'
                   f'margin:auto;padding:16px;max-width:1260px">{resumo}</body></html>')
    st.download_button("⬇ Baixar o resumo (abre no navegador, pronto para print)", pagina_html.encode("utf-8"),
                       file_name=f"resumo_time_{ano}.html", mime="text/html", key="vf_resumo")
    for a in list(avisos) + list(av_b or []):
        st.caption(f"⚠️ {a}")


# ── Conferência ──────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import sys
    import openpyxl
    falhas = 0

    def ok(nome, cond):
        global falhas
        falhas += not cond
        print(("ok    " if cond else "FALHA ") + nome)

    # A PLANILHA PELA CADEIA: um xlsx com as abas e colunas do Controle MS do
    # dono (15/10), lido por `controle_ms.ler` — o mesmo caminho da tela.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA_BASE
    ws.append(["MÊS", "MÊS/ANO", "Data", "Situação", "Plataforma", "Pedido", "SKU",
               "Produto", "Quantidade", "Custo Unitário", "FAT TOTAL", "DEVOLUÇÃO",
               "MARGEM C.", "VENDAS", "LPV", "UNI. CONT.", "FAT - DEV", "L.B"])
    # setembro fechado: UNI. CONT. = MARGEM C. ÷ LPV já calculada pela planilha
    ws.append(["set", "set2026", datetime(2026, 9, 10), "Entregue", "MercadoLivre1",
               "2000015134692443", "MS-1", "Argola", 1, 10.0, 100.0, 0, 41.26, 1, 20.63, 2.0, 100.0, 70.0])
    # pedido de dois produtos: duas linhas. A coluna VENDAS da planilha marca
    # 1 em TODA linha com faturamento e sem devolução (set/2026: 1.924 linhas,
    # 1.896 pedidos, VENDAS 1.857) — a venda é por linha, como lá.
    ws.append(["set", "set2026", datetime(2026, 9, 12), "Entregue", "Shopee1",
               "2609165EWUNJCD", "MS-2", "Xadrez A", 1, 5.0, 60.0, 0, 20.63, 1, 20.63, 1.0, 60.0, 40.0])
    ws.append(["set", "set2026", datetime(2026, 9, 12), "Entregue", "Shopee1",
               "2609165EWUNJCD", "MS-3", "Xadrez B", 1, 5.0, 40.0, 0, 0.0, 1, 20.63, 0.0, 40.0, 25.0])
    # devolução: não é venda, não tem UC
    ws.append(["set", "set2026", datetime(2026, 9, 15), "Devolvido", "MercadoLivre2",
               "2000099", "MS-4", "Pêndulo", 1, 8.0, 50.0, 50.0, 0.0, None, 20.63, None, 0.0, 0.0])
    # outubro em andamento: sem LPV nem UNI. CONT. ainda; uma venda sem SKU/custo
    ws.append(["out", "out2026", datetime(2026, 10, 5), "Entregue", "MercadoLivre1",
               "2000016000001", "MS-1", "Argola", 1, 10.0, 100.0, 0, 41.26, None, None, None, 100.0, 70.0])
    ws.append(["out", "out2026", datetime(2026, 10, 6), "A Enviar", "Shopee2",
               "2610071573P3B5", None, "Mini máquina", 1, None, 349.99, 0, 0.0, None, None, None, 349.99, 0.0])
    w2 = wb.create_sheet(ABA_CONSOLIDADO)
    w2.append([None, "FAT", 2024, 2025, 2026, None, "VENDAS", 2024, 2025, 2026])
    for i, m in enumerate(MESES):
        w2.append([None, m, 1000.0 + i, 9000.0 + i * 10, None, None, m, 10 + i, 90 + i, None])
    w2.append([None, "FAT", "%", "%", "%"])
    w2.append([None, "JAN", 0.5, 1.63, 0.1])     # o segundo bloco, de crescimento
    w3 = wb.create_sheet(ABA_DINAMICA)
    w3.append([None, 2026, None, 2025, None, None, "META DIA", 1, 2, None, "META 2026", 2700000])
    for q, (mf, mv) in enumerate(((594000, 7065.08), (648000, 7707.36), (692102.33, 8349.64), (788343.51, 9313.06)), 1):
        w3.append([None] * 15 + [f"{q}º Trimestre", 0.22, mf])
    w3.append([])
    for q, mv in enumerate((7065.08, 7707.36, 8349.64, 9313.06), 1):
        w3.append([None] * 15 + [f"{q}º Trimestre", 0.22, mv])
    buf = io.BytesIO()
    wb.save(buf)
    import controle_ms as _cms
    _g = _cms._bytes_do_arquivo
    try:
        _cms._bytes_do_arquivo = lambda: (buf.getvalue(), "")
        ler_planilha.clear()
        plan, av = ler_planilha()
    finally:
        _cms._bytes_do_arquivo = _g
        ler_planilha.clear()
    ok("as três abas chegam pelo leitor do Studio", not av and len(plan["base"]) == 6)
    d = dados_do_ano(2026, date(2026, 10, 15), plan)
    ok("metas da DINAMICA: o ano e os quatro trimestres, faturamento e vendas",
       d["metas"]["anual"] == 2700000 and d["metas"]["fat"][2] == 692102.33
       and d["metas"]["vendas"][3] == 9313.06 and d["metas"]["fonte"] == "DINAMICA")
    ok("2025 pelo FAT. CONSOLIDADO, e o bloco de % lá embaixo não sobrescreve",
       d["anterior"][1] == (9000.0, 90.0) and d["anterior"][12] == (9110.0, 101.0))
    st_ = d["meses"][9]
    ok("setembro: a devolução não é venda; VENDAS por linha, como a planilha",
       st_["vendas"] == 3 and st_["fat_liq"] == 200.0)
    ok("na conferência, o pedido de dois itens vira UM pedido",
       len(pedidos_da_planilha(plan["base"], 2026, 9)) == 3)
    ok("UC é margem de contribuição ÷ LPV (coluna da planilha), e não peças",
       round(st_["uc_total"], 2) == 3.0)
    ok("mês pronto só com SKU e custo em todas as vendas",
       pronto(d["meses"][9]) and not pronto(d["meses"][10]))
    ok("LPV de reserva é o do último mês que tem LPV (setembro)",
       round(d["lpv_reserva"], 2) == 20.63)
    o = d["meses"][10]
    ok("outubro sem LPV: UC pela mesma fórmula com o LPV de setembro",
       round(o["uc_total"], 2) == 2.0 and o["vendas"] == 2)
    ok("sem a DINAMICA do ano, a meta é o melhor ano anterior por trimestre",
       metas_da_dinamica(plan["dinamica"], 2027) is None
       and metas_do_melhor_ano(consolidado(plan["consolidado"]), 2027)["fat"][0] == 9000 + 9010 + 9020)
    ok("o número do pedido casa texto, número e lixo",
       chave_pedido(2.000015134692443e15) == "2000015134692443"
       and chave_pedido(" 2609165ewunjcd ") == "2609165EWUNJCD")

    # A CONFERÊNCIA PELA CADEIA: os pedidos do Bling no formato que a API
    # devolve, por `bling_api.pedido_resumido` de verdade.
    import bling_api as _bl
    canais = {204: "Mercado Livre MS", 305: "Shopee MS"}
    brutos = [{"numero": 101, "numeroLoja": "2000015134692443", "data": "2026-09-10", "total": 100.0,
               "loja": {"id": 204}, "situacao": {"id": 6}},
              {"numero": 102, "numeroLoja": "2609165EWUNJCD", "data": "2026-09-12", "total": 95.0,
               "loja": {"id": 305}, "situacao": {"id": 6}},
              {"numero": 103, "numeroLoja": "2609250SOBLING", "data": "2026-09-25", "total": 49.99,
               "loja": {"id": 305}, "situacao": {"id": 6}},
              {"numero": 104, "numeroLoja": "2610010OUTUBRO", "data": "2026-10-01", "total": 10.0,
               "loja": {"id": 305}, "situacao": {"id": 6}},
              {"numero": 105, "numeroLoja": "2609280CANCEL", "data": "2026-09-28", "total": 30.0,
               "loja": {"id": 204}, "situacao": {"id": 12}}]
    bling = [_bl.pedido_resumido(p, ids_cancelados={12}, canais=canais) for p in brutos]
    pl = pedidos_da_planilha(plan["base"], 2026, 9)
    cf = conferir(pl, bling, 2026, 9)
    ok("só na planilha: o pedido que o Bling não tem (com produto e SKU)",
       [x["pedido"] for x in cf["so_planilha"]] == ["2000099"])
    ok("só no Bling: só os do mês (o de 01/10 fica de fora), cancelado marcado",
       sorted(x["pedido"] for x in cf["so_bling"]) == ["2609250SOBLING", "2609280CANCEL"]
       and [x["cancelado"] for x in cf["so_bling"] if x["pedido"] == "2609280CANCEL"] == [True])
    ok("valor diferente: o pedido de dois itens soma 100 e o Bling diz 95",
       [(x["pedido"], x["diferenca"]) for x in cf["diferentes"]] == [("2609165EWUNJCD", 5.0)])
    ok("o canal vem com o nome do Bling",
       {x["canal"] for x in cf["so_bling"]} == {"Shopee MS", "Mercado Livre MS"})
    xb = planilha_da_conferencia(cf)
    _wb = openpyxl.load_workbook(io.BytesIO(xb))
    ok("o Excel das listas sai com as três abas",
       _wb.sheetnames == ["Só na planilha", "Só no Bling", "Valor diferente"]
       and _wb["Só no Bling"].max_row == 3)

    fm = realizado_por_mes(d, date(2026, 10, 15), 2026, 5000.0)
    ok("o mês corrente vem do Bling; os fechados, da planilha",
       fm[10] == 5000.0 and fm[9] == 200.0)
    import inspect
    ok("o Bling do mês é a mesma conta da Home",
       "_hg._faturamento_bling(hoje.year, hoje.month)" in inspect.getsource(bling_do_mes_corrente))
    # pedidos do Bling pela CADEIA: `pedidos_do_bling` → `bling_api.pedidos_resumidos`
    # → `pedido_resumido`, com só a rede trocada (formatos de `_paginar_pedidos`,
    # `id_de_cancelado` e `canais_de_venda`, conferidos em bling_api.py).
    import bling_api as _bl
    _orig = (_bl._paginar_pedidos, _bl.id_de_cancelado, _bl.canais_de_venda)
    _chamadas = []

    def _pag(conta, ini, fim, max_paginas=60):
        _chamadas.append(conta)
        if conta == "lg" and _chamadas.count("lg") == 1:
            return ([], "429 Too Many Requests")
        return ([{"numero": 7, "numeroLoja": f"PED-{conta}", "data": "2026-09-20 10:00:00",
                  "total": "49.90", "loja": {"id": 11}, "situacao": {"id": 9}}], "")
    try:
        _bl._paginar_pedidos = _pag
        _bl.id_de_cancelado = lambda c: ({9}, "")
        _bl.canais_de_venda = lambda c: ({11: "Shopee MS"}, "")
        _pedidos_da_conta.clear()
        p1, e1 = pedidos_do_bling(2026, 9)
        p2, e2 = pedidos_do_bling(2026, 9)
        pedidos_do_bling(2026, 9)
    finally:
        _bl._paginar_pedidos, _bl.id_de_cancelado, _bl.canais_de_venda = _orig
        _pedidos_da_conta.clear()
    ok("Bling pela cadeia: numeroLoja vira o pedido, canal com nome, cancelado marcado",
       p2 and p2[0]["pedido"] in ("PED-lg", "PED-ms") and p2[0]["canal"] == "Shopee MS"
       and p2[0]["cancelado"] and p2[0]["total"] == 49.9 and p2[0]["data"] == "2026-09-20")
    ok("falha do Bling NÃO fica guardada: o próximo clique tenta de novo; o que deu certo, sim",
       e1 == ["LG: 429 Too Many Requests"] and not e2 and len(p2) == 2
       and _chamadas.count("lg") == 2 and _chamadas.count("ms") == 1)
    _linha = lambda html, mes: next(t for t in html.split("<tr") if f">{mes}" in t)
    _tb = _tabela(d, fm, date(2026, 10, 15), 2026, False)
    ok("mês corrente sem CRESC. (mês parcial contra mês cheio dá −72% falso); o fechado tem",
       "▼" not in _linha(_tb, "OUT") and "▲" not in _linha(_tb, "OUT")
       and ("▼" in _linha(_tb, "SET") or "▲" in _linha(_tb, "SET")))
    ok("selo BLING só quando o valor do mês veio do Bling",
       "BLING" not in _linha(_tb, "OUT")
       and "BLING" in _linha(_tabela(d, fm, date(2026, 10, 15), 2026, True), "OUT"))
    tri = trimestres(d, fm)
    ok("trimestre só fecha com os três meses", tri[2][2] is None and tri[3][2] is None)
    h = html_resumo_time(d, tri, 5200.0, 4, 1.25, 1, 2700000, "OUT/2026")
    ok("o resumo para o time sai em uma linha (o markdown fecharia o HTML)", "\n" not in h)
    print(f"\nfalhas: {falhas}")
    sys.exit(1 if falhas else 0)
