"""visao_geral.py — o ano numa tela só: por finalidade em cima, por modalidade embaixo.

O PEDIDO, NAS PALAVRAS DO DONO (15/10)
--------------------------------------
"Eu preciso enxergar um macro parecido com o que tenho hoje na minha planilha
Controle MS, a aba DIN Finanças (...) para que eu não precise ir em cada aba
para pegar cada informação". Layout aprovado, com os ajustes dele:

    RESUMO POR FINALIDADE   mês a mês (JAN–DEZ), com LPV e C.O., botão R$ / %
                            e as linhas ANO e % ANO. "Mesmo havendo custo fixo
                            e assinaturas dentro do cartão de crédito, NO
                            RESUMO, o sistema tem que informar esse custo
                            nessa finalidade e não no cartão."
    DETALHE POR MODALIDADE  um módulo por modalidade (cartão, PIX, boleto…),
                            com as finalidades dentro, no mês escolhido.
    MERCADO LIVRE           publicidade, envios Full e outras cobranças, pagas
                            pelo saldo do Mercado Pago — fora das saídas.

DE ONDE VEM CADA NÚMERO: O STUDIO, E A PLANILHA ONDE FALTA EXTRATO
--------------------------------------------------------------------
"O que o sistema não tem importado em extrato pode configurar para ele
considerar o que está na minha planilha" — por BANCO e por MÊS (dono, 15/10):

    Inter LG (116183837) e Itaú LG (99659-5)   extrato no Studio naquele mês?
                                               Studio. Senão, SAIDAS PIX-BOLETO.
    Inter MS (167513915) e Itaú MS (99275-0)   só o Studio (a planilha não tem)
    Cartões                                    fatura no Studio naquele mês?
                                               Studio. Senão, aba CARTÕES.
    Impostos                                   imposto no Studio naquele mês?
                                               Studio. Senão, SIMPLES - FLEX.
    Mercado Livre                              aba ADS E CROSS (é dela o dado)

Nunca soma os dois lados do mesmo banco e mês: é o que contaria o mesmo PIX
duas vezes. Cada mês diz de onde veio (S = Studio, P = planilha).

Custo fixo, assinaturas, não operacional e folha de um mês FECHADO sem nenhum
pagamento lido entram pelo CADASTRO (C) — é o que a DIN Finanças faz, com o
custo fixo fixo em todo mês. Eles não entram nos módulos de modalidade: o
cadastro não diz por onde o dinheiro saiu.

UMA CONTA, E NÃO DUAS
---------------------
O pago do Studio é o de `resumo_mensal.itens_pagos` — a assinatura reconhecida
pelo cadastro, a fatura sem compras anexadas numa linha só. O cadastro é o de
`custo_real`, `nao_operacional` e `folha_salarial`. LPV e C.O. são os de
`lpv_mensal`. Nada aqui é uma segunda resposta.
"""
import re
import unicodedata
from datetime import date

import streamlit as st

MESES = ("JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT",
         "NOV", "DEZ")
_MES_TXT = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set",
            "out", "nov", "dez")

MODALIDADES = ("CARTÃO", "PIX", "CHEQUE", "BOLETO", "SISPAG", "DÉBITO EM CONTA",
               "TRANSFERÊNCIA", "NÃO IDENTIFICADO")
ROTULO_MOD = {"SISPAG": "SISPAG (BOLETO/TED/DÉBITO)"}

# As contas que a planilha cobre, pelo número (dono, 15/10). As MS ficam de
# fora: a SAIDAS PIX-BOLETO só traz INTER LG e ITAU LG.
BANCOS_LG = {"INTER LG": ("inter-", "116183837"),
             "ITAU LG": ("itau-", "996595")}

# A ordem das colunas do resumo. Finalidade que não está aqui entra depois,
# pela ordem do valor — o resumo mostra o que existe, e não uma lista fixa.
ORDEM_FIN = ("CUSTO FIXO", "ASSINATURAS", "NÃO OPERACIONAL", "FOLHA",
             "MERCADORIA", "SERVIÇOS", "CONSUMO INTERNO", "IMPOSTOS", "ADS",
             "CHEQUES")
_FIN_IGUAL = {"SERVICO": "SERVIÇOS", "SERVICOS": "SERVIÇOS",
              "IMPOSTO": "IMPOSTOS", "DASS": "IMPOSTOS", "DARF": "IMPOSTOS",
              "DAS": "IMPOSTOS", "NAO OPERACIONAL": "NÃO OPERACIONAL",
              "FATURA DO CARTAO": "CARTÃO SEM DETALHE",
              "": "SEM FINALIDADE"}

# O que não diz o que é vai para o fim da tabela, e não para o meio.
_NO_FIM = ("OUTROS", "CARTÃO SEM DETALHE", "SEM FINALIDADE")

# Os cheques saem do Itaú LG (as tarifas de cheque da SAIDAS PIX-BOLETO são
# todas de lá). Mês em que o Itaú LG vem da planilha não tem os cheques — a
# aba é de PIX e boleto —, e eles entram pelo cadastro de cheques do Studio.
BANCO_DOS_CHEQUES = "ITAU LG"

ML_GRUPOS = ("PUBLICIDADE", "ENVIOS FULL", "OUTRAS COBRANÇAS")


def _chave(t):
    t = unicodedata.normalize("NFKD", str(t or ""))
    return " ".join("".join(c for c in t if not unicodedata.combining(c))
                    .upper().split())


def _num(v):
    try:
        f = float(v)
        return f if f == f else 0.0
    except (TypeError, ValueError):
        try:
            return float(str(v).replace(".", "").replace(",", "."))
        except (TypeError, ValueError):
            return 0.0


def finalidade(fin):
    """A finalidade como o resumo a mostra: SERVIÇO e SERVIÇOS são uma só."""
    k = _chave(fin)
    if k in _FIN_IGUAL:
        return _FIN_IGUAL[k]
    return str(fin or "").strip().upper()


def modalidade(texto, cartao=False):
    """Por onde o dinheiro saiu, pela descrição do extrato. Pura.

    A ORDEM IMPORTA: "TAR CHEQUE EMITIDO" é tarifa (débito), e não cheque;
    "Debito titulo" do Inter é boleto, e não débito em conta.
    """
    if cartao:
        return "CARTÃO"
    t = _chave(texto)
    if re.search(r"\bTAR\b|TARIFA|\bIOF\b|JUROS|ENCARGO|PARCELA|DEBITO AUT", t):
        return "DÉBITO EM CONTA"
    if "PIX" in t:
        return "PIX"
    if "SISPAG" in t or "FORNECEDORES" in t:
        return "SISPAG"
    if "BOLETO" in t or "TITULO" in t:
        return "BOLETO"
    if "CHEQUE" in t or "COMPENSAD" in t:
        return "CHEQUE"
    if re.search(r"\bTED\b|\bTEF\b|\bDOC\b|TRANSF", t):
        return "TRANSFERÊNCIA"
    if "DEBITO" in t:
        return "DÉBITO EM CONTA"
    return "NÃO IDENTIFICADO"


def banco_da_conta(conta):
    """'INTER LG' / 'ITAU LG' para as contas que a planilha cobre; '' no resto."""
    c = str(conta or "").strip().lower()
    dig = re.sub(r"\D", "", c)
    for banco, (pref, num) in BANCOS_LG.items():
        if c.startswith(pref) and num in dig:
            return banco
    return ""


def banco_da_planilha(txt):
    t = _chave(txt)
    if t in ("TAU LG", "ITAU LG"):          # 6 linhas da planilha vêm "TAU LG"
        return "ITAU LG"
    return t


def mes_da_planilha(v):
    """'jan2026' / data → 'AAAA-MM'. '' quando não dá."""
    if isinstance(v, date):
        return v.strftime("%Y-%m")
    t = str(v or "").strip().lower()
    m = re.match(r"([a-zç]{3})\D*(\d{4})", t)
    if m and m.group(1) in _MES_TXT:
        return f"{m.group(2)}-{_MES_TXT.index(m.group(1)) + 1:02d}"
    m = re.match(r"(\d{4})-(\d{2})", t)
    return f"{m.group(1)}-{m.group(2)}" if m else ""


def grupo_ml(descricao):
    k = _chave(descricao)
    if "PUBLICIDADE" in k:
        return "PUBLICIDADE"
    if "FULL" in k:
        return "ENVIOS FULL"
    return "OUTRAS COBRANÇAS"


# ── O ANO, ITEM A ITEM ───────────────────────────────────────────────────────

def itens_do_studio(itens_por_mes):
    """[{mes, modalidade, finalidade, valor, fonte, banco, cartao}] do Studio.

    `itens_por_mes` = {'AAAA-MM': resumo_mensal.itens_pagos(..., com_origem=True)}.
    """
    fora = []
    for mes, itens in (itens_por_mes or {}).items():
        for i in itens:
            fin = ("ASSINATURAS" if i.get("quadro") == "Assinaturas"
                   else finalidade(i.get("finalidade")))
            cart = bool(i.get("cartao"))
            fora.append({"mes": mes, "fonte": "S", "valor": float(i["valor"]),
                         "finalidade": fin, "cartao": cart,
                         "modalidade": modalidade(i.get("descricao"), cart),
                         "banco": "" if cart else banco_da_conta(i.get("conta"))})
    return fora


def juntar(ano, studio, saidas, cartoes, impostos, cheques=None):
    """Os itens do ano: o Studio, e a planilha onde o Studio não tem. Pura.

    `saidas`, `cartoes`, `impostos` = linhas das abas SAIDAS PIX-BOLETO 2026,
    CARTÕES e SIMPLES - FLEX, como `controle_ms.ler` as devolve (dicionário
    por linha, com os nomes das colunas de lá). `cheques` = `cheques.carregar`.
    """
    pref = f"{int(ano):04d}-"
    st_ano = [i for i in (studio or []) if i["mes"].startswith(pref)]
    tem_banco = {(i["mes"], i["banco"]) for i in st_ano if i["banco"]}
    tem_cartao = {i["mes"] for i in st_ano
                  if i["cartao"] and i["finalidade"] != "CARTÃO SEM DETALHE"}
    tem_imposto = {i["mes"] for i in st_ano if i["finalidade"] == "IMPOSTOS"}
    fora = list(st_ano)
    for r in (saidas or []):
        mes = mes_da_planilha(r.get("Data"))
        banco = banco_da_planilha(r.get("Banco"))
        if (not mes.startswith(pref) or banco not in BANCOS_LG
                or (mes, banco) in tem_banco):
            continue
        v = -_num(r.get("Valor"))                    # saída vem negativa
        if v:
            fora.append({"mes": mes, "fonte": "P", "valor": v,
                         "finalidade": finalidade(r.get("FINALIDADE")),
                         "modalidade": modalidade(r.get("Operação")),
                         "banco": banco, "cartao": False})
    cart_planilha = set()
    for r in (cartoes or []):
        mes = mes_da_planilha(r.get("MÊS/ANO"))
        if not mes.startswith(pref) or mes in tem_cartao:
            continue
        v = _num(r.get("Valor"))
        if v:
            cart_planilha.add(mes)
            fora.append({"mes": mes, "fonte": "P", "valor": v,
                         "finalidade": finalidade(r.get("FINALIDADE")),
                         "modalidade": "CARTÃO", "banco": "", "cartao": True})
    # A fatura paga sem compras anexadas é o MESMO dinheiro que a aba CARTÕES
    # detalha: com a planilha no mês, a linha sem detalhe sai.
    fora = [i for i in fora if not (i["fonte"] == "S" and i["mes"] in cart_planilha
                                    and i["finalidade"] == "CARTÃO SEM DETALHE")]
    import cheques as _ch
    for c in (cheques or []):
        if not _ch.eh_tipo(c, "CHEQUE") or _ch.situacao_de(c.get("situacao")) != "DEBITADO":
            continue
        mes = _ch.texto_data(c.get("vencimento"))[:7]
        if not mes.startswith(pref) or (mes, BANCO_DOS_CHEQUES) in tem_banco:
            continue
        v = _num(c.get("valor"))
        if v:
            fora.append({"mes": mes, "fonte": "S", "valor": v,
                         "finalidade": "CHEQUES", "modalidade": "CHEQUE",
                         "banco": BANCO_DOS_CHEQUES, "cartao": False})
    for r in (impostos or []):
        if finalidade(r.get("FINALIDADE")) != "IMPOSTOS":
            continue
        mes = mes_da_planilha(r.get("MÊS/ANO"))
        if not mes.startswith(pref) or mes in tem_imposto:
            continue
        v = _num(r.get("VALOR"))
        if v:
            # DASS e DARF são pagos como guia, isto é, boleto (dono, 15/10).
            fora.append({"mes": mes, "fonte": "P", "valor": v,
                         "finalidade": "IMPOSTOS", "modalidade": "BOLETO",
                         "banco": "", "cartao": False})
    return fora


def com_cadastro(itens, ano, hoje, cadastro_de):
    """Mês fechado sem pagamento lido de custo fixo, assinaturas, não
    operacional ou folha recebe o CADASTRO (fonte C). Pura.

    `cadastro_de(ano, mes)` → {finalidade: valor} do cadastro do Studio.
    """
    fora = list(itens or [])
    for m in range(1, 13):
        if (int(ano), m) >= (hoje.year, hoje.month):
            break
        mes = f"{int(ano):04d}-{m:02d}"
        pagos = {i["finalidade"] for i in fora if i["mes"] == mes}
        for fin, v in (cadastro_de(int(ano), m) or {}).items():
            if fin not in pagos and v:
                fora.append({"mes": mes, "fonte": "C", "valor": float(v),
                             "finalidade": fin, "modalidade": "", "banco": "",
                             "cartao": False})
    return fora


def colunas(itens):
    """As finalidades do resumo: as de ORDEM_FIN que existem, e o resto."""
    soma = {}
    for i in itens or []:
        soma[i["finalidade"]] = soma.get(i["finalidade"], 0.0) + i["valor"]
    fixas = [f for f in ORDEM_FIN if f in soma]
    resto = sorted((f for f in soma if f not in ORDEM_FIN + _NO_FIM),
                   key=lambda f: -soma[f])
    return fixas + resto + [f for f in _NO_FIM if f in soma]


def resumo(itens, ano):
    """{mes_idx: {finalidade: valor}} e {mes_idx: fontes}. Pura."""
    valores, fontes = {}, {}
    for i in itens or []:
        if not i["mes"].startswith(f"{int(ano):04d}-"):
            continue
        m = int(i["mes"][5:7])
        d = valores.setdefault(m, {})
        d[i["finalidade"]] = round(d.get(i["finalidade"], 0.0) + i["valor"], 2)
        fontes.setdefault(m, set()).add(i["fonte"])
    return valores, fontes


def modulos(itens, mes):
    """{modalidade: {finalidade: valor}} do mês 'AAAA-MM'. Cadastro fica fora."""
    fora = {}
    for i in itens or []:
        if i["mes"] != mes or i["fonte"] == "C":
            continue
        d = fora.setdefault(i["modalidade"], {})
        d[i["finalidade"]] = round(d.get(i["finalidade"], 0.0) + i["valor"], 2)
    return {m: fora[m] for m in MODALIDADES if m in fora}


def mercado_livre(linhas, ano):
    """{mes_idx: {grupo: valor}} da aba ADS E CROSS. Pura."""
    fora = {}
    for r in linhas or []:
        mes = mes_da_planilha(r.get("MÊS/ANO"))
        if not mes.startswith(f"{int(ano):04d}-"):
            continue
        v = _num(r.get("COBRANÇA"))
        if not v:
            continue
        d = fora.setdefault(int(mes[5:7]), {})
        g = grupo_ml(r.get("DESCRIÇÃO"))
        d[g] = round(d.get(g, 0.0) + v, 2)
    return fora


# ── Leitura ──────────────────────────────────────────────────────────────────

ABAS_PLANILHA = ("SAIDAS PIX-BOLETO 2026", "CARTÕES", "SIMPLES - FLEX",
                 "ADS E CROSS")


@st.cache_data(ttl=600, show_spinner=False)
def da_planilha():
    """({aba: [linhas]}, [avisos]). Uma leitura de cada aba a cada 10 min."""
    import controle_ms as _cms
    fora, avisos = {}, []
    for aba in ABAS_PLANILHA:
        df, erro = _cms.ler(aba)
        if erro:
            avisos.append(f"{aba}: {erro}")
            fora[aba] = []
            continue
        fora[aba] = df.to_dict("records") if df is not None else []
    return fora, avisos


def _cadastro_do_mes(grade, ajustes, assin, mapa, nao_op, folha):
    import custo_real as _cr
    import nao_operacional as _no
    import folha_salarial as _fs

    def _f(ano, mes):
        return {"CUSTO FIXO": _cr.total_custo_fixo(grade, ajustes, ano, mes, mapa),
                "ASSINATURAS": _cr.total_assinaturas(assin, ano, mes, mapa),
                "NÃO OPERACIONAL": float(_no.total_no_mes(nao_op, ano, mes) or 0.0),
                "FOLHA": float(_fs.total_da_folha(folha, ano, mes, ajustes) or 0.0)}
    return _f


@st.cache_data(ttl=120, show_spinner=False)
def dados_do_ano(ano, hoje):
    """(itens, ml, lpv, co, avisos) do ano. Tudo o que a tela desenha.

    CACHE DE 120 s, o mesmo de `lancamentos.carregar`: cada clique no botão
    R$ / % e no mês do detalhe redesenha a página, e reaplicar o cadastro de
    favorecidos ao ano inteiro a cada clique seria trabalho jogado fora.
    """
    import lancamentos as _lan
    import assinaturas as _as
    import resumo_mensal as _rm
    avisos = []
    pref = f"{int(ano):04d}-"
    lancs = _lan.aplicar_cadastro(
        [l for l in _lan.carregar() if _lan.mes_de(l).startswith(pref)])
    assin = _as.carregar() or []
    nomes = _rm.nomes_das_assinaturas(assin)
    por_mes = {}
    for l in lancs:
        por_mes.setdefault(_lan.mes_de(l), []).append(l)
    studio = itens_do_studio({m: _rm.itens_pagos(ls, nomes, com_origem=True)
                              for m, ls in por_mes.items()})
    plan, av = da_planilha()
    avisos += av
    try:
        import cheques as _ch
        chs = _ch.carregar()
    except Exception as e:
        chs = []
        avisos.append(f"cheques: {type(e).__name__}")
    itens = juntar(ano, studio, plan.get("SAIDAS PIX-BOLETO 2026"),
                   plan.get("CARTÕES"), plan.get("SIMPLES - FLEX"), chs)
    try:
        import ajustes as _aj
        import custo_fixo as _cf
        import custo_real as _cr
        import nao_operacional as _no
        import folha_salarial as _fs
        cad = _cadastro_do_mes(_cf.carregar("custo_fixo"), _aj.carregar(), assin,
                               _cr.carregar(), _no.carregar(), _fs.carregar())
        itens = com_cadastro(itens, ano, hoje, cad)
    except Exception as e:
        avisos.append(f"cadastro do custo fixo: {type(e).__name__}")
    ml = mercado_livre(plan.get("ADS E CROSS"), ano)
    lpv, co = {}, {}
    try:
        import lpv_mensal as _lm
        _l, _e1 = _lm.lpv_do_ano(int(ano))
        _c, _e2 = _lm.custo_operacional_do_ano(int(ano))
        lpv = {m: v.get("lpv") for m, v in (_l or {}).items()}
        co = {m: (v.get("co") if not v.get("motivo") else None)
              for m, v in (_c or {}).items()}
        avisos += [e for e in (_e1, _e2) if e]
    except Exception as e:
        avisos.append(f"LPV e C.O.: {type(e).__name__}")
    return itens, ml, lpv, co, avisos


def saidas_restantes(hoje):
    """O que ainda vai sair no mês corrente: provisionado − pago, por quadro
    (`resumo_mensal.montar`, a mesma conta do Resumo Mensal). (valor, erros)."""
    import lancamentos as _lan
    import assinaturas as _as
    import custo_real as _cr
    import previsto as _pv
    import resumo_mensal as _rm
    lancs = _lan.do_mes(hoje.year, hoje.month)
    assin = _as.carregar() or []
    prev, erros = _pv.do_mes(hoje.year, hoje.month)
    q = _rm.montar(_rm.itens_pagos(lancs, _rm.nomes_das_assinaturas(assin)),
                   prev, _cr.total_assinaturas(assin, hoje.year, hoje.month,
                                               _cr.carregar()), False)
    falta = sum(max((v["provisionado"] or 0.0) - v["pago"], 0.0)
                for v in q.values() if v["provisionado"] is not None)
    return round(falta, 2), erros


# ── Tela ─────────────────────────────────────────────────────────────────────

FUNDO, CARTAO, BORDA, TILE = "#3c3c3c", "#333333", "#4a4a4a", "#2c2c2c"
TEXTO, SEC = "#E8EAED", "#9AA3AE"
VERDE, AMBAR, AZUL, VERMELHO, LILAS = ("#4CC38A", "#E0A13A", "#7FB8F0",
                                       "#E5534B", "#B9A2FF")
COR_MOD = {"CARTÃO": LILAS, "PIX": AZUL, "CHEQUE": AMBAR, "BOLETO": VERMELHO,
           "SISPAG": TEXTO, "DÉBITO EM CONTA": VERDE, "TRANSFERÊNCIA": AZUL,
           "NÃO IDENTIFICADO": SEC}


def _brl(v, casas=0):
    if v is None:
        return "—"
    s = f"{float(v):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def _pct(v):
    return f"{v:.1f}%".replace(".", ",")


def _caixa(titulo, cor, miolo, borda=BORDA):
    return (f'<div style="background:{CARTAO};border:1px solid {borda};'
            f'border-radius:12px;padding:12px;margin-top:12px;">'
            f'<div style="text-align:center;font-size:12px;font-weight:800;'
            f'letter-spacing:.1em;color:{cor};margin-bottom:8px;">{titulo}</div>'
            f'{miolo}</div>')


def _selo(fontes):
    fora = ""
    for f, cor, fundo in (("P", AMBAR, "#5a4a2a"), ("C", AZUL, "#2a3b52"),
                          ("S", VERDE, "#24452f")):
        if f in (fontes or ()):
            fora += (f' <span style="font-size:9px;padding:1px 5px;'
                     f'border-radius:6px;background:{fundo};color:{cor};">{f}</span>')
    return fora


def html_resumo(valores, fontes, cols, lpv, co, ano, em_pct=False):
    """A tabela do resumo, em R$ ou em %. Uma linha só (markdown)."""
    th = (lambda c: f'<th style="text-align:right !important;padding:6px 7px;color:{SEC};'
          f'font-size:11px;font-weight:700;border-bottom:1px solid {BORDA};'
          f'white-space:nowrap;">{c}</th>')
    td = (lambda v, c=TEXTO, b=False: f'<td style="text-align:right;padding:5px 7px;'
          f'font-size:12px;white-space:nowrap;color:{c};'
          f'{"font-weight:800;" if b else ""}">{v}</td>')
    cab = ("<tr><th></th>" + "".join(th(c) for c in cols) + th("TOTAL")
           + th("LPV") + th("C.O.") + "</tr>")
    corpo = ""
    for m in range(1, 13):
        lin = valores.get(m) or {}
        tot = round(sum(lin.values()), 2)
        cel = "".join(
            td(("—" if not lin.get(c) else
                (_pct(lin[c] / tot * 100) if em_pct and tot else _brl(lin[c]))))
            for c in cols)
        corpo += (f'<tr><td style="padding:5px 7px;font-weight:700;'
                  f'white-space:nowrap;">{MESES[m - 1]}{_selo(fontes.get(m))}</td>'
                  + cel + td(_brl(tot) if lin else "—", VERDE, True)
                  + td(_brl(lpv.get(m), 2), AMBAR) + td(_brl(co.get(m)), AMBAR)
                  + "</tr>")
    soma = {c: round(sum((valores.get(m) or {}).get(c, 0.0) for m in valores), 2)
            for c in cols}
    T = round(sum(soma.values()), 2)
    corpo += (f'<tr style="border-top:2px solid {BORDA};"><td style="padding:6px 7px;'
              f'font-weight:800;">ANO</td>'
              + "".join(td(_brl(soma[c]), TEXTO, True) for c in cols)
              + td(_brl(T), VERDE, True) + td("") + td("") + "</tr>")
    corpo += (f'<tr><td style="padding:6px 7px;font-weight:800;color:{SEC};">% ANO</td>'
              + "".join(td(_pct(soma[c] / T * 100) if T else "—", SEC, True)
                        for c in cols)
              + td("100%" if T else "—", SEC, True) + td("") + td("") + "</tr>")
    legenda = (f'<div style="font-size:11px;color:{SEC};margin-top:6px;">'
               f'Cada gasto entra na finalidade dele, venha de onde vier. '
               f'<b style="color:{VERDE};">S</b> extratos do Studio · '
               f'<b style="color:{AMBAR};">P</b> sua planilha (banco e mês sem '
               f'extrato) · <b style="color:{AZUL};">C</b> cadastro do Studio '
               f'(mês fechado sem pagamento lido)</div>')
    return _caixa(f"RESUMO POR FINALIDADE · {ano}", VERDE,
                  '<div style="overflow-x:auto;"><table style="border-collapse:'
                  f'collapse;width:100%;">{cab}{corpo}</table></div>' + legenda)


def _modulo(nome, cor, itens):
    tot = round(sum(itens.values()), 2)
    linhas = "".join(
        f'<div style="display:flex;justify-content:space-between;gap:8px;'
        f'padding:5px 0;border-bottom:1px solid {BORDA};font-size:13px;">'
        f'<span>{f.capitalize()}</span><b>{_brl(v)}'
        f'<span style="color:{SEC};font-weight:500;font-size:11px;"> · '
        f'{_pct(v / tot * 100) if tot else "—"}</span></b></div>'
        for f, v in sorted(itens.items(), key=lambda kv: -kv[1]))
    return (f'<div style="background:{TILE};border-top:3px solid {cor};'
            f'border-radius:10px;padding:10px 12px;">'
            f'<div style="text-align:center;font-weight:800;color:{cor};'
            f'font-size:13px;letter-spacing:.06em;margin-bottom:6px;">{nome}</div>'
            f'{linhas}<div style="display:flex;justify-content:space-between;'
            f'padding-top:7px;font-weight:800;color:{cor};"><span>Total</span>'
            f'<span>{_brl(tot)}</span></div></div>')


def html_modulos(mods, ml_mes):
    cards = [_modulo(ROTULO_MOD.get(m, m), COR_MOD.get(m, TEXTO), v)
             for m, v in mods.items() if v]
    if ml_mes:
        cards.append(_modulo("MERCADO LIVRE · SALDO MP", LILAS,
                             {g.capitalize(): ml_mes.get(g, 0.0)
                              for g in ML_GRUPOS if ml_mes.get(g)}))
    if not cards:
        return f'<div style="text-align:center;color:{SEC};">Nada lido neste mês.</div>'
    return ('<div style="display:grid;grid-template-columns:repeat(auto-fit,'
            f'minmax(260px,1fr));gap:10px;">{"".join(cards)}</div>')


def html_mes(restantes):
    def tile(a, b):
        return (f'<div style="background:{TILE};border-radius:8px;padding:8px 10px;'
                f'text-align:center;"><div style="font-size:11px;color:{SEC};">{a}'
                f'</div><div style="font-weight:800;">{b}</div></div>')
    return ('<div style="display:grid;grid-template-columns:repeat(3,1fr);gap:8px;">'
            + tile("Saídas restantes", _brl(restantes))
            + tile("A receber", "em espera") + tile("Balanço", "em espera")
            + "</div>")


ABAS = ("Financeira", "Faturamento")


def pagina(usuario_logado=None):
    from datetime import datetime
    import auth as _auth
    import placar_core as _pc
    # A trava é da página também, e não só da aba: a URL leva aqui sem a aba.
    if not _auth.eh_dono(usuario_logado):
        st.error("A Visão Geral é só do dono.")
        return
    hoje = datetime.now(_pc.FUSO).date()
    c1, c2 = st.columns([4, 1])
    c1.markdown("### 🧭 Visão Geral")
    ano = c2.selectbox("Ano", list(range(hoje.year, 2023, -1)), index=0,
                       key="vg_ano")
    # AS DUAS ABAS (dono, 15/10): "fará 2 abas, financeira e faturamento".
    aba = st.radio("Visão", ABAS, horizontal=True, key="vg_aba",
                   label_visibility="collapsed")
    if aba == ABAS[1]:
        import visao_faturamento as _vf
        _vf.pagina(usuario_logado, int(ano), hoje)
        return
    with st.spinner("Lendo extratos, cadastro e planilha…"):
        itens, ml, lpv, co, avisos = dados_do_ano(int(ano), hoje)
    if int(ano) == hoje.year:
        try:
            rest, err_r = saidas_restantes(hoje)
        except Exception as e:
            rest, err_r = None, [type(e).__name__]
        st.markdown(_caixa(f"ESTE MÊS · {MESES[hoje.month - 1]}", AMBAR,
                           html_mes(rest), AMBAR), unsafe_allow_html=True)
        st.caption("A receber e Balanço ficam em espera até a integração com "
                   "os marketplaces.")
        if err_r:
            st.caption("⚠️ Saídas restantes sem: " + " · ".join(map(str, err_r)))
    modo = st.radio("Mostrar", ["R$", "%"], horizontal=True, key="vg_modo",
                    label_visibility="collapsed")
    valores, fontes = resumo(itens, ano)
    st.markdown(html_resumo(valores, fontes, colunas(itens), lpv, co, ano,
                            em_pct=(modo == "%")), unsafe_allow_html=True)
    st.markdown("&nbsp;", unsafe_allow_html=True)
    padrao = hoje.month - 1 if int(ano) == hoje.year else 11
    mes_txt = st.radio("Mês do detalhe", list(MESES), index=padrao,
                       horizontal=True, key="vg_mes",
                       label_visibility="collapsed")
    m = MESES.index(mes_txt) + 1
    st.markdown(_caixa(f"DETALHE POR MODALIDADE · {mes_txt}", AZUL,
                       html_modulos(modulos(itens, f"{int(ano):04d}-{m:02d}"),
                                    ml.get(m))),
                unsafe_allow_html=True)
    st.caption("O Mercado Livre é pago pelo saldo do Mercado Pago e não entra "
               "nas saídas do resumo. Custo fixo pelo cadastro (C) não entra "
               "nos módulos: o cadastro não diz por onde o dinheiro saiu.")
    for a in avisos:
        st.warning(a)


# ── Conferência ──────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import io
    import sys
    falhas = 0

    def ok(nome, cond):
        global falhas
        falhas += not cond
        print(("ok    " if cond else "FALHA ") + nome)

    # AS DESCRIÇÕES DE VERDADE dos extratos do dono (Itaú 99275-0 de 13/10 e
    # Inter 167513915 de 09/10) e as operações da SAIDAS PIX-BOLETO.
    _casos = {"BOLETO PAGO": "BOLETO", "SISPAG FORNECEDORES": "SISPAG",
              "PAGAMENTOS A FORNECEDORES": "SISPAG",
              "PARCELA GIRO 18/48": "DÉBITO EM CONTA", "IOF": "DÉBITO EM CONTA",
              "JUROS LIMITE": "DÉBITO EM CONTA",
              'Pix enviado: "Cp :60701190-LITTLE GLASS"': "PIX",
              'Debito titulo KG: "Ref.: 325003255 - 16"': "BOLETO",
              "TAR CHEQUE EMITIDO": "DÉBITO EM CONTA", "PIX ENVIADO": "PIX",
              "TAR SAQUE BCO 24HS": "DÉBITO EM CONTA"}
    _err = {k: modalidade(k) for k, v in _casos.items() if modalidade(k) != v}
    ok("a modalidade sai da descrição real dos extratos", not _err)
    if _err:
        print("      ", _err)
    ok("as contas LG pelo número; as MS ficam fora da planilha",
       banco_da_conta("itau-0099659-5") == "ITAU LG"
       and banco_da_conta("inter-116183837") == "INTER LG"
       and banco_da_conta("itau-0099275-0") == ""
       and banco_da_conta("inter-167513915") == "")
    ok("'TAU LG' da planilha é o Itaú LG; 'jan2026' é 2026-01",
       banco_da_planilha("TAU LG") == "ITAU LG"
       and mes_da_planilha("jan2026") == "2026-01"
       and mes_da_planilha(date(2026, 8, 3)) == "2026-08")

    # A CADEIA DO STUDIO: lançamentos na forma de `lancamentos.carregar`,
    # a fatura pelo leitor de PDF e por `preparar_fatura`, e os itens por
    # `resumo_mensal.itens_pagos` de verdade.
    import fatura_pdf as _fp
    import lancamentos as _lan
    import resumo_mensal as _rm
    _fat, _ = _fp.lancamentos_do_texto(
        "02/09   ANTHROPIC* CLAUDE SUB      R$ 550,00\n"
        "05/09   UBER TRIP                  R$ 40,00\n", ano_padrao=2026)
    _cart = _lan.preparar_fatura(_fat, "2026-09")
    for _c in _cart:
        _c["conta"] = _lan.CONTA_CARTAO + "fatura.pdf"
        _c["finalidade"] = "CUSTO FIXO" if "CLAUDE" in _c["descricao"] else "SERVIÇOS"
    _cc = [{"data": "2026-08-10", "valor": -1000.0, "conta": "inter-116183837",
            "descricao": 'Pix enviado: "Cp :1-FORNECEDOR"', "favorecido": "FORNECEDOR",
            "finalidade": "MERCADORIA", "tipo": ""},
           {"data": "2026-08-12", "valor": -300.0, "conta": "itau-0099275-0",
            "descricao": "BOLETO PAGO", "favorecido": "ENEL",
            "finalidade": "CUSTO FIXO", "tipo": ""}]
    _assin = [{"item": "Claude", "favorecido": "CLAUDE", "valor_mensal": 550.0}]
    _nomes = _rm.nomes_das_assinaturas(_assin)
    _pm = {}
    for _l in _cc + _cart:
        _pm.setdefault(_lan.mes_de(_l), []).append(_l)
    _studio = itens_do_studio({m: _rm.itens_pagos(ls, _nomes, com_origem=True)
                               for m, ls in _pm.items()})

    # A PLANILHA PELA CADEIA: um xlsx com as colunas de lá, lido por
    # `controle_ms.ler` — o mesmo caminho de `da_planilha`.
    import openpyxl as _ox
    import controle_ms as _cms
    _wb = _ox.Workbook()
    _ws = _wb.active
    _ws.title = "SAIDAS PIX-BOLETO 2026"
    _ws.append([None, "Data", "Operação", "DESTINATÁRIO", "FINALIDADE", "Valor", "Banco"])
    _ws.append([None, date(2026, 8, 5), "PIX ENVIADO", "X", "MERCADORIA", -999.0, "INTER LG"])
    _ws.append([None, date(2026, 8, 6), "SISPAG FORNECEDORES", "Y", "OUTROS", -200.0, "ITAU LG"])
    _ws.append([None, date(2026, 7, 6), "PIX ENVIADO", "Z", "EMBALAGEM", -50.0, "INTER LG"])
    _w2 = _wb.create_sheet("CARTÕES")
    _w2.append(["MÊS/ANO", "Compra", "VENCIMENTO", "DESTINATÁRIO", "FINALIDADE", "Valor", "Banco"])
    _w2.append(["set2026", None, None, "UBER", "SERVIÇOS", 40.0, "Inter"])
    _w2.append(["ago2026", None, None, "ESTAC", "ESTACIONAMENTO", 70.0, "Inter"])
    _w3 = _wb.create_sheet("SIMPLES - FLEX")
    _w3.append([None])
    _w3.append(["MÊS/ANO", "DATA", "VALOR", "EMPRESA", "FINALIDADE"])
    _w3.append(["ago2026", None, 945.73, "MARTINSOUSA", "DASS"])
    _w3.append(["ago2026", None, 10.0, "TM", None])
    _w4 = _wb.create_sheet("ADS E CROSS")
    _w4.append(["MÊS/ANO", "FECHAMENTO", "VENCIMENTO", "DESCRIÇÃO", "FINALIDADE", "COBRANÇA", "CONTA"])
    _w4.append(["ago2026", None, None, "PUBLICIDADE", "SERVIÇOS", 3119.76, "ML1"])
    _w4.append(["ago2026", None, None, "ENVIOS FULL", "SERVIÇOS", 39.09, "ML1"])
    _w4.append(["ago2026", None, None, "PÁGINA DO ML", "SERVIÇOS", 99.0, "ML1"])
    _buf = io.BytesIO()
    _wb.save(_buf)
    _g = _cms._bytes_do_arquivo
    try:
        _cms._bytes_do_arquivo = lambda: (_buf.getvalue(), "")
        da_planilha.clear()
        _plan, _av = da_planilha()
    finally:
        _cms._bytes_do_arquivo = _g
        da_planilha.clear()
    ok("as quatro abas da planilha chegam pelo leitor do Studio",
       not _av and len(_plan["SAIDAS PIX-BOLETO 2026"]) == 3
       and len(_plan["SIMPLES - FLEX"]) == 2)

    _it = juntar(2026, _studio, _plan["SAIDAS PIX-BOLETO 2026"],
                 _plan["CARTÕES"], _plan["SIMPLES - FLEX"])
    _ago = [i for i in _it if i["mes"] == "2026-08"]
    ok("Inter LG de agosto com extrato: o PIX da planilha NÃO entra (nunca soma)",
       not any(i["fonte"] == "P" and i["banco"] == "INTER LG" for i in _ago)
       and any(i["fonte"] == "S" and i["valor"] == 1000.0 for i in _ago))
    ok("Itaú LG de agosto sem extrato: vem da planilha, SISPAG",
       any(i["fonte"] == "P" and i["banco"] == "ITAU LG"
           and i["modalidade"] == "SISPAG" and i["valor"] == 200.0 for i in _ago))
    ok("julho sem extrato do Inter LG: a planilha preenche",
       any(i["mes"] == "2026-07" and i["fonte"] == "P"
           and i["finalidade"] == "EMBALAGEM" for i in _it))
    ok("a conta MS vem sempre do Studio",
       any(i["mes"] == "2026-08" and i["modalidade"] == "BOLETO"
           and i["finalidade"] == "CUSTO FIXO" for i in _it))
    _set = [i for i in _it if i["mes"] == "2026-09"]
    ok("setembro com fatura no Studio: o cartão da planilha NÃO entra",
       not any(i["fonte"] == "P" and i["cartao"] for i in _set))
    ok("agosto sem fatura no Studio: o cartão vem da planilha",
       any(i["fonte"] == "P" and i["cartao"] and i["valor"] == 70.0 for i in _ago))
    ok("imposto da SIMPLES - FLEX: só DASS/DARF, no mês sem imposto no Studio, como boleto",
       [(round(i["valor"], 2), i["modalidade"]) for i in _ago
        if i["finalidade"] == "IMPOSTOS"] == [(945.73, "BOLETO")])

    import cheques as _ch_t
    _chq = [_ch_t.normalizar({"tipo": "CHEQUE", "vencimento": "2026-08-15",
                              "valor": 4000.0, "situacao": "DEBITADO"}),
            _ch_t.normalizar({"tipo": "CHEQUE", "vencimento": "2026-08-20",
                              "valor": 500.0, "situacao": "EM ABERTO"})]
    _it_c = juntar(2026, _studio, [], [], [], _chq)
    ok("Itaú LG de agosto sem extrato: o cheque compensado entra pelo cadastro",
       [(i["valor"], i["modalidade"]) for i in _it_c
        if i["finalidade"] == "CHEQUES"] == [(4000.0, "CHEQUE")])
    _st_itau = _studio + [{"mes": "2026-08", "fonte": "S", "valor": 1.0,
                           "finalidade": "OUTROS", "modalidade": "PIX",
                           "banco": "ITAU LG", "cartao": False}]
    ok("com extrato do Itaú LG no mês, o cheque já está nele: não soma de novo",
       not any(i["finalidade"] == "CHEQUES"
               for i in juntar(2026, _st_itau, [], [], [], _chq)))
    ok("OUTROS vai para o fim da tabela",
       colunas([{"finalidade": "OUTROS", "valor": 9e9},
                {"finalidade": "EMBALAGEM", "valor": 1.0},
                {"finalidade": "CUSTO FIXO", "valor": 1.0}])
       == ["CUSTO FIXO", "EMBALAGEM", "OUTROS"])

    _val, _fon = resumo(_it, 2026)
    ok("NO RESUMO a assinatura paga no cartão é ASSINATURAS, e não cartão",
       _val[9].get("ASSINATURAS") == 550.0 and "CARTÃO" not in _val[9])
    _mods = modulos(_it, "2026-09")
    ok("NO DETALHE o cartão abre por finalidade, com a assinatura dentro",
       _mods.get("CARTÃO") == {"ASSINATURAS": 550.0, "SERVIÇOS": 40.0})
    ok("cada mês diz de onde veio",
       _fon[8] == {"S", "P"} and _fon[9] == {"S"} and _fon[7] == {"P"})

    _cad = com_cadastro(_it, 2026, date(2026, 10, 15),
                        lambda a, m: {"CUSTO FIXO": 22842.0, "FOLHA": 40000.0})
    _jul = {i["finalidade"]: (i["fonte"], i["valor"]) for i in _cad
            if i["mes"] == "2026-07" and i["fonte"] == "C"}
    ok("mês fechado sem pagamento lido: custo fixo e folha pelo cadastro (C)",
       _jul == {"CUSTO FIXO": ("C", 22842.0), "FOLHA": ("C", 40000.0)})
    ok("agosto pagou custo fixo: o cadastro NÃO entra por cima",
       not any(i["mes"] == "2026-08" and i["fonte"] == "C"
               and i["finalidade"] == "CUSTO FIXO" for i in _cad))
    ok("o mês corrente nunca recebe cadastro",
       not any(i["mes"] == "2026-10" and i["fonte"] == "C" for i in _cad))
    ok("cadastro não entra nos módulos de modalidade",
       "" not in modulos(_cad, "2026-07"))

    _ml = mercado_livre(_plan["ADS E CROSS"], 2026)
    ok("Mercado Livre: publicidade, envios Full e o resto separados",
       _ml[8] == {"PUBLICIDADE": 3119.76, "ENVIOS FULL": 39.09,
                  "OUTRAS COBRANÇAS": 99.0})

    _v2, _f2 = resumo(_cad, 2026)
    _h = html_resumo(_v2, _f2, colunas(_cad), {9: 20.63}, {8: 18975.32}, 2026)
    _hp = html_resumo(_v2, _f2, colunas(_cad), {}, {}, 2026, em_pct=True)
    ok("a tabela: meses em maiúscula, ANO e % ANO, LPV e C.O.",
       all(m in _h for m in MESES) and ">ANO<" in _h and "% ANO" in _h
       and "R$ 20,63" in _h and "R$ 18.975" in _h)
    ok("o botão % troca os valores por porcentagem do mês",
       "%" in _hp.split("% ANO")[0] and "100%" in _hp)
    ok("o HTML não tem quebra de linha (o markdown fecharia a tabela)",
       "\n" not in _h and "\n" not in html_modulos(_mods, _ml[8]))
    ok("Mercado Livre entra no detalhe como módulo próprio",
       "MERCADO LIVRE · SALDO MP" in html_modulos(_mods, _ml[8]))

    # 19/10 (dono): a Visão Geral saiu da barra de Gestão e virou a aba do
    # meio da Home — Painel | Visão Geral | Histórico —, só para o dono.
    import gestao as _ge_t
    import inspect as _insp_t
    ok("a Visão Geral é a aba do meio da Home, e só o dono a vê",
       _ge_t.HOME_ABAS[1] == "🧭 Visão Geral"
       and "_vg.pagina(usuario_logado)" in _insp_t.getsource(_ge_t.pagina_home)
       and "eh_dono" in _insp_t.getsource(_ge_t.abas_da_home))
    import inspect as _insp
    ok("a página tranca por conta própria",
       "_auth.eh_dono(usuario_logado)" in _insp.getsource(pagina))

    print(f"\nfalhas: {falhas}")
    sys.exit(1 if falhas else 0)
