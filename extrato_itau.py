"""extrato_itau.py — o extrato do Itaú lido como lançamento, não como texto.

O QUE ELE RESOLVE
-----------------
O extrato chega em `.xlsx` e hoje é conferido a olho: o dono abre, procura o
que é custo fixo, tira do que vai para a planilha, e o que sobra ele digita.
São duas horas de trabalho por mês em algo que o arquivo já diz.

Aqui ele vira lista de lançamentos com TIPO, e o tipo é o que permite a
conferência automática: `CH COMPENSADO 001 000504` é o cheque 504 dando baixa,
`DA ELETROPAULO` é a luz do mês, `BUSINESS 6202-5907` é a fatura do cartão.

O FORMATO, CONFERIDO NO ARQUIVO DE VERDADE
------------------------------------------
Extrato de 01/08 a 31/08 de 2026, conta 0099659-5 (LITTLE GLASS), 149 linhas:

    linha 1-9    cabeçalho: Atualização, Nome, Agência, Conta, Período
    linha 10     Data | Lançamento | Razão Social | CPF/CNPJ | Valor (R$)
    linha 11+    os lançamentos

Duas linhas NÃO são lançamento e precisam sair, senão viram gasto de valor
vazio: `SALDO ANTERIOR` e `SALDO TOTAL DISPONÍVEL DIA` — esta última aparece
uma vez por dia, 20 vezes no mês.

O SINAL MANDA
-------------
Valor negativo é saída, positivo é entrada. Não se deduz isso da descrição:
`SISPAG MAST CD` tanto paga quanto recebe, e chutar pelo nome inverteria o
fluxo de caixa inteiro.
"""

import io
import re

# As linhas que o extrato imprime e que não são dinheiro se movendo.
NAO_E_LANCAMENTO = ("SALDO ANTERIOR", "SALDO TOTAL DISPON", "SALDO DO DIA")

# Cada tipo é uma pergunta que o Studio sabe responder sozinho depois.
# A ordem importa: a primeira que casar vence, e as mais específicas vêm antes.
TIPOS = [
    ("cheque",        (r"^CH\s+COMPENSADO",)),
    ("tarifa",        (r"^TAR\b", r"^TARIFA")),
    ("fatura_cartao", (r"^BUSINESS\b",)),
    ("folha",         (r"^SISPAG\s+SALARIOS",)),
    ("debito_auto",   (r"^DA\s", r"^SABESP", r"^ELETROPAULO", r"^ENEL",
                       r"^COMGAS", r"^VIVO", r"^CLARO", r"^TIM\b")),
    ("emprestimo",    (r"^PARCELA\s+GIRO", r"CAPITAL\s+DE\s+GIRO", r"^EMPREST")),
    ("aplicacao",     (r"^APLICACAO", r"^RESGATE", r"^RENDIMENTOS")),
    ("boleto",        (r"^BOLETO\s+PAGO",)),
    ("pix",           (r"^PIX\s",)),
    ("ted",           (r"^TED\s",)),
    ("sispag",        (r"^SISPAG",)),
]


# ── Quem é cada cartão ───────────────────────────────────────────────────────
#
# O extrato bancário NÃO traz o número do cartão: "BUSINESS 6202-5907" é o
# contrato, e nenhum dígito dele aparece na fatura. A ligação foi feita
# cruzando data de vencimento com ordem de grandeza, e confirmada pelo dono em
# 17/09/2026:
#
#   6202-5907  pago dia 03/08, R$ 16.684,36  -> VISA INFINITE ...3417
#   4005-0759  pago dia 03/08, R$  1.226,74  -> MASTERCARD ...9113
#   4004-7311  pago dia 11/08, R$  1.623,63  -> MASTERCARD ...3312 (vence dia 11)
#
# Fica no código porque é identidade de contrato bancário, não cadastro que
# muda: cartão novo entra aqui, e a tela avisa quando aparecer um contrato
# desconhecido em vez de somar tudo num balaio só.
CARTOES = {
    "6202-5907": "VISA INFINITE ...3417",
    "4005-0759": "MASTERCARD ...9113",
    "4004-7311": "MASTERCARD ...3312",
}


def qual_cartao(descricao):
    """O cartão de um lançamento BUSINESS. "" quando não é fatura.

    Devolve o contrato cru quando ele não está em CARTOES: cartão novo precisa
    aparecer como desconhecido, e não ser somado em silêncio com os outros.
    """
    import re as _re
    t = str(descricao or "").strip().upper()
    if not t.startswith("BUSINESS"):
        return ""
    m = _re.search(r"(\d{4})[-\s]?(\d{4})", t)
    if not m:
        return ""
    contrato = f"{m.group(1)}-{m.group(2)}"
    return CARTOES.get(contrato, f"contrato {contrato} (não cadastrado)")


def _num(v):
    """O valor como float. None quando a célula está vazia."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        _f = float(v)
        # `nan` E float: sem esta linha ele atravessa o isinstance
        # inteiro e vai parar na planilha, que recusa gravar — "Out of
        # range float values are not JSON compliant: nan". Foi assim
        # que a importacao de 604 cheques morreu depois de ler tudo.
        if _f != _f or _f in (float('inf'), float('-inf')):
            return None          # aqui o vazio é None, e não zero
        return _f
    t = str(v).strip().replace("R$", "").replace(" ", "")
    if not t:
        return None
    # O arquivo vem com ponto decimal ("-968.6"), mas extrato exportado de
    # outra tela vem com vírgula. Aceitar os dois sai mais barato que
    # descobrir na hora errada que o valor virou mil vezes maior.
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _data(v):
    """'03/08/2026' -> date. None quando não dá para ler."""
    from datetime import date, datetime
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()[:10]
    for f in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(t, f).date()
        except ValueError:
            continue
    return None


def classificar(descricao):
    """O tipo do lançamento. "outro" quando nenhum padrão casa.

    "outro" não é falha: é o que o Studio vai perguntar ao dono uma vez, para
    nunca mais perguntar. Chutar um tipo seria pior — classificação errada não
    se percebe, ela só aparece como indicador torto no fim do mês.
    """
    t = str(descricao or "").strip().upper()
    if not t:
        return "outro"
    for nome, padroes in TIPOS:
        for p in padroes:
            if re.search(p, t):
                return nome
    return "outro"


def numero_do_cheque(descricao):
    """O número do cheque em 'CH COMPENSADO 001 000504' -> '504'.

    Sem os zeros à esquerda: no canhoto e na planilha ele é 504, e comparar
    '000504' com '504' daria duas respostas para o mesmo cheque.
    """
    t = str(descricao or "").strip().upper()
    if not t.startswith("CH COMPENSADO"):
        return ""
    numeros = re.findall(r"\d+", t)
    if not numeros:
        return ""
    return str(int(numeros[-1]))


def ler(dados):
    """(lançamentos, cabeçalho, erro) do `.xlsx` do Itaú.

    Cada lançamento: {"data", "descricao", "razao_social", "cnpj", "valor",
    "tipo", "cheque"}. `valor` negativo é saída.
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True,
                                    data_only=True)
    except Exception as e:
        return [], {}, f"Não consegui abrir o arquivo: {str(e)[:140]}"
    try:
        ws = wb[wb.sheetnames[0]]
        cabecalho, linhas, achou_titulo = {}, [], False
        for linha in ws.iter_rows(values_only=True):
            celulas = [("" if c is None else str(c).strip()) for c in linha]
            if not any(celulas):
                continue
            primeira = celulas[0]
            if not achou_titulo:
                # O cabeçalho vem como pares rótulo/valor antes da tabela.
                if primeira.upper().startswith("DATA"):
                    achou_titulo = True
                elif len(celulas) > 1 and primeira.endswith(":"):
                    cabecalho[primeira.rstrip(":").lower()] = celulas[1]
                continue
            if any(primeira.upper().startswith(x) or
                   celulas[1].upper().startswith(x)
                   for x in NAO_E_LANCAMENTO):
                continue
            valor = _num(celulas[4] if len(celulas) > 4 else None)
            data = _data(primeira)
            if valor is None or data is None:
                continue
            desc = celulas[1] if len(celulas) > 1 else ""
            linhas.append({
                "data": data,
                "descricao": desc,
                "razao_social": celulas[2] if len(celulas) > 2 else "",
                "cnpj": celulas[3] if len(celulas) > 3 else "",
                "valor": valor,
                "tipo": classificar(desc),
                "cheque": numero_do_cheque(desc),
                "cartao": qual_cartao(desc),
            })
        return linhas, cabecalho, ""
    except Exception as e:
        return [], {}, f"Falhou lendo o extrato: {str(e)[:140]}"
    finally:
        wb.close()


def resumo(lancamentos):
    """{tipo: {"n": int, "entrada": float, "saida": float}} — para a tela."""
    fora = {}
    for l in (lancamentos or []):
        d = fora.setdefault(l["tipo"], {"n": 0, "entrada": 0.0, "saida": 0.0})
        d["n"] += 1
        if l["valor"] < 0:
            d["saida"] += -l["valor"]
        else:
            d["entrada"] += l["valor"]
    return fora


# ── Conferência ──────────────────────────────────────────────────────────────
# As linhas abaixo têm a forma do extrato real — conferida no arquivo de agosto
# de 2026 —, com valores inventados. Extrato de verdade não entra no
# repositório.
if __name__ == "__main__":
    falhas = 0

    def ok(nome, cond):
        global falhas
        falhas += not cond
        print(("ok    " if cond else "FALHA ") + nome)

    ok("cheque compensado e reconhecido",
       classificar("CH COMPENSADO 001 000504") == "cheque")
    ok("o numero do cheque sai sem zeros a esquerda",
       numero_do_cheque("CH COMPENSADO 001 000504") == "504")
    ok("quem nao e cheque nao devolve numero",
       numero_do_cheque("PIX ENVIADO") == ""
       and numero_do_cheque("") == "")

    ok("debito automatico da luz", classificar("DA  ELETROPAULO") == "debito_auto")
    ok("agua tambem", classificar("SABESP") == "debito_auto")
    ok("telefone tambem", classificar("DA  VIVO FIXO") == "debito_auto")
    ok("fatura do cartao", classificar("BUSINESS      6202-5907") == "fatura_cartao")
    ok("folha", classificar("SISPAG SALARIOS") == "folha")
    ok("emprestimo", classificar("PARCELA GIRO 12/48") == "emprestimo")
    ok("emprestimo pela outra redacao",
       classificar("EMPREST CAPITAL DE GIRO") == "emprestimo")
    ok("boleto", classificar("BOLETO PAGO TM LOGISTICA") == "boleto")
    ok("pix", classificar("PIX ENVIADO") == "pix")
    ok("ted", classificar("TED RECEBIDA 745.0001.MODA M") == "ted")
    ok("tarifa", classificar("TAR CHEQUE EMITIDO") == "tarifa")
    ok("rendimento e aplicacao", classificar("RENDIMENTOS REND PAGO") == "aplicacao")
    # SISPAG generico fica DEPOIS de SISPAG SALARIOS: a ordem e o que separa
    # folha de repasse de cartao.
    ok("sispag generico nao vira folha",
       classificar("SISPAG MAST CD") == "sispag")
    ok("o que nao casa vira 'outro', e nao um chute",
       classificar("COISA QUE NAO EXISTE") == "outro"
       and classificar("") == "outro" and classificar(None) == "outro")

    ok("cada contrato BUSINESS vira o cartao dele",
       qual_cartao("BUSINESS      6202-5907") == "VISA INFINITE ...3417"
       and qual_cartao("BUSINESS      4005-0759") == "MASTERCARD ...9113"
       and qual_cartao("BUSINESS      4004-7311") == "MASTERCARD ...3312")
    ok("cartao novo aparece como desconhecido, e nao somado no balaio",
       "não cadastrado" in qual_cartao("BUSINESS      9999-0000"))
    ok("o que nao e fatura nao devolve cartao",
       qual_cartao("PIX ENVIADO") == "" and qual_cartao("") == "")

    ok("valor com ponto decimal", _num("-968.6") == -968.6)
    ok("valor com virgula e milhar", _num("-1.301,61") == -1301.61)
    ok("celula vazia nao vira zero", _num("") is None and _num(None) is None)

    from datetime import date as _d
    ok("data brasileira", _data("03/08/2026") == _d(2026, 8, 3))
    ok("data invalida nao derruba", _data("banana") is None)

    _l = [{"tipo": "cheque", "valor": -968.6},
          {"tipo": "cheque", "valor": -1301.61},
          {"tipo": "pix", "valor": 5866.16}]
    _r = resumo(_l)
    ok("o resumo separa entrada de saida",
       _r["cheque"]["n"] == 2 and abs(_r["cheque"]["saida"] - 2270.21) < 0.01
       and _r["cheque"]["entrada"] == 0.0
       and abs(_r["pix"]["entrada"] - 5866.16) < 0.01)

    print("\nfalhas:", falhas)
